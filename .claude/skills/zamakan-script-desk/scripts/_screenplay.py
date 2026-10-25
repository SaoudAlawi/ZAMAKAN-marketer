"""Shared screenplay layout for the script desk: page setup, Word styles, Arabic/English text runs and PDF export.

Layout follows the standard screenplay page (Courier 12 pt, 1.5" binding margin, 1" elsewhere, character cue at
3.7", dialogue 2.5"-6", parenthetical 3.1"-5.6" from the binding edge). Arabic pages mirror it (binding margin on the
right, text right to left), as in WriterDuet's right-to-left template. Word and LibreOffice treat a paragraph's
"left" indent and alignment as its start side in right-to-left paragraphs, so the same indents work for both.
"""
import json
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn  # noqa: F401 (re-exported)
from docx.shared import Inches, Pt, RGBColor

SKILL = Path(__file__).resolve().parents[1]
FONTS = SKILL / "assets" / "fonts"
VOCAB = json.loads((SKILL / "references" / "vocabulary.json").read_text(encoding="utf-8"))
STRUCTURES = json.loads((SKILL / "references" / "structures.json").read_text(encoding="utf-8"))

LATIN_FONT = "Courier Prime"
ARABIC_FONT = "Noto Naskh Arabic"
PAPER = {"a4": (8.27, 11.69), "letter": (8.5, 11.0)}
# element: (start indent, end indent, alignment, space before in lines) in inches from the text margins
ELEMENTS = {
    "scene_heading": (0.0, 0.0, "start", 1),
    "action": (0.0, 0.0, "start", 1),
    "character": (2.2, 0.0, "start", 1),
    "parenthetical": (1.6, 1.9, "start", 0),
    "dialogue": (1.0, 1.5, "start", 0),
    "transition": (0.0, 0.0, "end", 1),
    "centered": (0.0, 0.0, "center", 1),
    "section": (0.0, 0.0, "start", 1),
}
STYLE_NAMES = {"scene_heading": "Scene Heading", "action": "Action", "character": "Character", "parenthetical": "Parenthetical",
               "dialogue": "Dialogue", "transition": "Transition", "centered": "Centered", "section": "Section"}
STYLE_TO_TYPE = {v.lower(): k for k, v in STYLE_NAMES.items()}
STYLE_TO_TYPE.update({"general": "action", "shot": "scene_heading", "heading": "scene_heading"})
AR_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")
ARABIC_RE = re.compile(r"[؀-ۿ]")


def is_arabic(text):
    letters = re.findall(r"[A-Za-z؀-ۿ]", text)
    return bool(letters) and sum(bool(ARABIC_RE.match(c)) for c in letters) / len(letters) > 0.5


def digits(n, lang, arabic_digits=True):
    s = str(n)
    return s.translate(AR_DIGITS) if lang == "ar" and arabic_digits else s


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), str(v))
    return e


def set_run_fonts(rpr, latin=LATIN_FONT, arabic=ARABIC_FONT, size=12, size_cs=14, rtl=False, bold=False):
    for old in rpr.findall(qn("w:rFonts")):
        rpr.remove(old)
    rpr.insert(0, _el("w:rFonts", **{"w:ascii": latin, "w:hAnsi": latin, "w:cs": arabic, "w:eastAsia": latin}))
    if bold:
        rpr.append(_el("w:b"))
        rpr.append(_el("w:bCs"))
    rpr.append(_el("w:sz", **{"w:val": int(size * 2)}))
    rpr.append(_el("w:szCs", **{"w:val": int(size_cs * 2)}))
    if rtl:
        rpr.append(_el("w:rtl"))


def new_document(lang="ar", paper=None, arabic_font=ARABIC_FONT, arabic_size=14):
    """A blank screenplay document with page setup and the element styles (Scene Heading, Action, Character…)."""
    paper = paper or ("a4" if lang == "ar" else "letter")
    doc = Document()
    sec = doc.sections[0]
    w, h = PAPER[paper]
    sec.page_width, sec.page_height = Inches(w), Inches(h)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = Inches(0.5)
    # binding margin is on the reading-start side: right for Arabic, left for English
    sec.left_margin, sec.right_margin = (Inches(1.0), Inches(1.5)) if lang == "ar" else (Inches(1.5), Inches(1.0))
    if lang == "ar":
        sec._sectPr.append(_el("w:bidi"))
    rtl = lang == "ar"
    normal = doc.styles["Normal"]
    set_run_fonts(normal.element.get_or_add_rPr(), arabic=arabic_font, size=12, size_cs=arabic_size, rtl=rtl)
    npf = normal.paragraph_format
    npf.space_before = npf.space_after = Pt(0)
    npf.line_spacing = 1.0 if rtl else Pt(12)
    if rtl:
        normal.element.get_or_add_pPr().append(_el("w:bidi"))
    for kind, name in STYLE_NAMES.items():
        st = doc.styles.add_style(name, 1)
        st.base_style = normal
        st.quick_style = True
        start, end, align, before = ELEMENTS[kind]
        pf = st.paragraph_format
        pf.left_indent, pf.right_indent = Inches(start), Inches(end)
        pf.alignment = {"start": WD_ALIGN_PARAGRAPH.LEFT, "end": WD_ALIGN_PARAGRAPH.RIGHT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
        pf.space_before = Pt(12 * before if not rtl else 14 * before)
        pf.keep_with_next = kind in ("scene_heading", "character", "parenthetical")
        if rtl:
            st.element.get_or_add_pPr().append(_el("w:bidi"))
        bold = kind in ("scene_heading", "character") and rtl or kind == "section"
        set_run_fonts(st.element.get_or_add_rPr(), arabic=arabic_font, size=12, size_cs=arabic_size, rtl=rtl, bold=bold)
        if kind == "section":
            st.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        st = doc.styles[name]
        if rtl:
            st.element.get_or_add_pPr().append(_el("w:bidi"))
        rpr = st.element.get_or_add_rPr()
        for old in rpr.findall(qn("w:rFonts")):
            rpr.remove(old)
        rpr.insert(0, _el("w:rFonts", **{"w:cs": arabic_font}))
        if rtl:
            rpr.append(_el("w:rtl"))
    return doc


def add_page_numbers(doc, lang):
    """Page number at the top, on the end side (top-right in English, top-left in Arabic), from page 2."""
    sec = doc.sections[-1]
    sec.different_first_page_header_footer = True
    p = sec.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if lang == "ar":
        p._p.get_or_add_pPr().append(_el("w:bidi"))
    run = p.add_run()
    for tag, extra in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", {}), ("w:fldChar", {"w:fldCharType": "end"})):
        e = _el(tag, **extra)
        if tag == "w:instrText":
            e.set(qn("xml:space"), "preserve")
            e.text = " PAGE "
        run._r.append(e)
    if lang != "ar":
        p.add_run(".")


def add_element(doc, kind, text, lang, scene_no=None, arabic_digits=True):
    """One screenplay element as a paragraph in its style; Arabic runs are marked right-to-left."""
    p = doc.add_paragraph(style=STYLE_NAMES[kind])
    if kind == "scene_heading" and scene_no is not None:
        text = f"{digits(scene_no, lang, arabic_digits)}. {text}"
    if lang == "en" and kind in ("scene_heading", "character", "transition"):
        text = text.upper()
    run = p.add_run(text)
    if lang == "ar":
        run._r.get_or_add_rPr().append(_el("w:rtl"))
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def to_pdf(docx_path, out_dir=None):
    """Convert with LibreOffice, using the bundled fonts (Courier Prime, Noto Naskh Arabic) so the PDF embeds them."""
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        return None, "LibreOffice not found: the Word file is ready; install LibreOffice for a PDF"
    docx_path = Path(docx_path)
    out_dir = Path(out_dir or docx_path.parent)
    with tempfile.TemporaryDirectory() as home:
        fonts = Path(home) / ".fonts"
        fonts.mkdir()
        for f in FONTS.glob("*.ttf"):
            shutil.copy(f, fonts / f.name)
        env = {**os.environ, "HOME": home}
        subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
                       env=env, capture_output=True, timeout=180)
    pdf = out_dir / (docx_path.stem + ".pdf")
    return (pdf, None) if pdf.exists() else (None, "PDF conversion failed")


def pdf_pages(pdf):
    try:
        import pymupdf
        with pymupdf.open(pdf) as d:
            return d.page_count
    except Exception:  # noqa: BLE001
        return None
