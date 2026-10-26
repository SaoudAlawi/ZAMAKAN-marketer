#!/usr/bin/env python3
"""Documentary transcripts → one bites index, so Claude can find and pick what people actually said without reading
hours of interviews, and every quote stays verbatim with its timecode.

Reads subtitle files (.srt, .vtt), timestamped text (.txt) and Word (.docx) transcripts, from people or transcription
services. Lines it understands:
  [00:01:23] Salem: text      00:01:23 Salem: text      Salem (00:01:23): text
  Salem  00:01:23             (a header line; the text follows on the next lines)
  Salem: text                 (no timecode: kept, but flagged)
A Word transcript can also be a table: timecode | speaker | text.
Short subtitle cues by the same speaker are joined into one bite until a sentence ends or the speaker pauses.
Bracketed directions «(يضحك)» stay out of the text and go to "direction".

Usage:
  python transcripts.py ingest a.srt b.txt [--out bites.json] [--speaker NAME] [--interviewer NAME] [--map "Speaker 1=Salem"]
  python transcripts.py index bites.json [--speaker NAME] [--grep WORD] [--width 90]   # one line per bite
  python transcripts.py show bites.json A004 A010-A014                                # full text of chosen bites
"""
import argparse
import json
import re
import sys
from pathlib import Path

DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
TS = r"(?:\d{1,2}:)?\d{1,2}:\d{2}(?:[.,:]\d{1,3})?"
SPEAKER = r"[^\W\d][\w .'\-]{0,30}?"  # a short name, starting with a letter
END = re.compile(r"[.!?؟…»\"]\s*$")
DIRECTION = re.compile(r"[\(\[]([^\)\]]{1,60})[\)\]]")
PATTERNS = [
    re.compile(rf"^\[?(?P<ts>{TS})\]?\s*[-–]?\s*(?P<sp>{SPEAKER})\s*:\s*(?P<tx>.*)$"),  # [ts] Name: text
    re.compile(rf"^(?P<sp>{SPEAKER})\s*[\(\[](?P<ts>{TS})[\)\]]\s*:?\s*(?P<tx>.*)$"),  # Name (ts): text
    re.compile(rf"^(?P<sp>{SPEAKER})\s+[\(\[]?(?P<ts>{TS})[\)\]]?\s*$"),  # header: Name ts
    re.compile(rf"^\[?(?P<ts>{TS})\]?\s*$"),  # header: ts alone
    re.compile(rf"^\[?(?P<ts>{TS})\]?\s+(?P<tx>.+)$"),  # [ts] text
    re.compile(rf"^(?P<sp>{SPEAKER})\s*:\s*(?P<tx>.+)$"),  # Name: text
]


def secs(ts):
    """'01:02:03,500' / '02:03' / '1:02:03:12' (frames ignored) → seconds."""
    ts = ts.translate(DIGITS).replace(",", ".")
    parts = ts.split(":")
    if len(parts) == 4:  # HH:MM:SS:FF
        parts = parts[:3]
    total = 0.0
    for p in parts:
        total = total * 60 + float(p)
    return round(total, 3)


def tc(s):
    if s is None:
        return ""
    s = int(s)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def split_direction(text):
    dirs = DIRECTION.findall(text)
    clean = re.sub(r"\s{2,}", " ", DIRECTION.sub("", text)).strip()
    return clean, dirs


def read_subtitles(path):
    """SRT or VTT → raw cues [(in, out, text)]."""
    raw = path.read_text(encoding="utf-8-sig").translate(DIGITS).replace("\r", "")
    cues = []
    for block in re.split(r"\n\s*\n", raw):
        lines = [x for x in block.split("\n") if x.strip()]
        for i, ln in enumerate(lines):
            m = re.match(rf"^\s*({TS})\s*-->\s*({TS})", ln)
            if m:
                text = " ".join(x.strip() for x in lines[i + 1:])
                text = re.sub(r"<[^>]+>", "", text)  # vtt styling
                if text:
                    cues.append((secs(m.group(1)), secs(m.group(2)), text))
                break
    return cues


def read_lines(lines):
    """Timestamped text lines → raw cues [(in, None, text-with-speaker)]."""
    cues, cur = [], None
    for ln in lines:
        ln = ln.translate(DIGITS).strip()
        if not ln:
            continue
        for k, pat in enumerate(PATTERNS):
            m = pat.match(ln)
            if not m:
                continue
            g = m.groupdict()
            if k in (2, 3):  # a header: text follows
                cur = [secs(g["ts"]), g.get("sp"), []]
                cues.append(cur)
            else:
                ts = secs(g["ts"]) if g.get("ts") else None
                cur = [ts, g.get("sp"), [g["tx"]]]
                cues.append(cur)
            break
        else:
            if cur is None:
                cur = [None, None, []]
                cues.append(cur)
            cur[2].append(ln)
    return [(c[0], None, (f"{c[1]}: " if c[1] else "") + " ".join(c[2])) for c in cues if c[2]]


def read_docx(path):
    from docx import Document
    d = Document(path)
    for t in d.tables:  # timecode | speaker | text
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        good = [r for r in rows if len(r) >= 3 and re.fullmatch(TS, r[0].translate(DIGITS))]
        if len(good) >= max(2, len(rows) // 2):
            return [(secs(r[0]), None, f"{r[1]}: {r[2]}" if r[1] else r[2]) for r in good]
    return read_lines(p.text for p in d.paragraphs)


def to_bites(cues, letter, file, default_speaker=None, names=None):
    """Raw cues → bites: speaker resolved and inherited, short cues by one speaker joined into sentences."""
    names = names or {}
    bites, speaker, warn = [], default_speaker, 0
    for t_in, t_out, text in cues:
        m = re.match(rf"^\s*[-–]?\s*({SPEAKER})\s*:\s*(.*)$", text)
        new_speaker = None
        if m and len(m.group(1).split()) <= 3:
            new_speaker, text = names.get(m.group(1).strip(), m.group(1).strip()), m.group(2)
        elif re.match(r"^\s*[-–]\s*", text):  # "- text": a new (unnamed) speaker
            text = re.sub(r"^\s*[-–]\s*", "", text)
        text, dirs = split_direction(text)
        prev = bites[-1] if bites else None
        joined = (prev and new_speaker is None and prev["out"] is not None and t_in is not None and t_in - prev["out"] <= 1.0
                  and not END.search(prev["text"]) and len(prev["text"].split()) < 60)
        if joined:
            prev["text"] = (prev["text"] + " " + text).strip()
            prev["out"] = t_out
            prev["direction"] += dirs
            continue
        speaker = new_speaker or speaker
        if t_in is None:
            warn += 1
        bites.append({"id": f"{letter}{len(bites) + 1:03d}", "file": file, "speaker": speaker or "?", "in": t_in, "out": t_out,
                      "text": text, "direction": dirs})
    for i, b in enumerate(bites):  # lines without an out point end where the next one starts
        if b["out"] is None and b["in"] is not None:
            nxt = next((x["in"] for x in bites[i + 1:] if x["in"] is not None), None)
            b["out"] = nxt if nxt is not None else b["in"] + max(2, len(b["text"].split()) / 2.5)
    for b in bites:
        b["tc_in"], b["tc_out"] = tc(b["in"]), tc(b["out"])
        if not b["direction"]:
            b.pop("direction")
    bites = [b for b in bites if b["text"] or b.get("direction")]  # a silence with a direction is a moment too
    return bites, warn


def ingest(paths, default_speaker=None, interviewer=None, names=None):
    out = {"schema": "zamakan-bites/1", "files": [], "bites": []}
    for i, p in enumerate(paths):
        p = Path(p)
        letter = chr(ord("A") + i) if i < 26 else f"F{i}"
        ext = p.suffix.lower()
        cues = read_subtitles(p) if ext in (".srt", ".vtt") else read_docx(p) if ext == ".docx" else \
            read_lines(p.read_text(encoding="utf-8-sig").splitlines())
        bites, warn = to_bites(cues, letter, p.name, default_speaker, names)
        for b in bites:
            if interviewer and b["speaker"] == interviewer:
                b["interviewer"] = True
        speakers = sorted({b["speaker"] for b in bites})
        out["files"].append({"letter": letter, "file": p.name, "bites": len(bites), "speakers": speakers,
                             "duration": tc(max((b["out"] or 0) for b in bites) if bites else 0), "no_timecode": warn})
        out["bites"] += bites
    return out


def expand(ids, bites):
    order = [b["id"] for b in bites]
    got = []
    for token in ids:
        for part in token.split(","):
            if "-" in part:
                a, z = part.split("-")
                if a in order and z in order:
                    got += order[order.index(a):order.index(z) + 1]
            elif part:
                got.append(part)
    return got


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a1 = sub.add_parser("ingest")
    a1.add_argument("files", nargs="+")
    a1.add_argument("--out", default="bites.json")
    a1.add_argument("--speaker", help="speaker for lines with no name")
    a1.add_argument("--interviewer", help="the person asking (their lines are marked, so they are not picked as bites)")
    a1.add_argument("--map", action="append", default=[], help='rename a speaker: "Speaker 1=Salem"')
    a2 = sub.add_parser("index")
    a2.add_argument("bites")
    a2.add_argument("--speaker")
    a2.add_argument("--grep")
    a2.add_argument("--width", type=int, default=90)
    a3 = sub.add_parser("show")
    a3.add_argument("bites")
    a3.add_argument("ids", nargs="+")
    a = ap.parse_args()
    if a.cmd == "ingest":
        names = dict(x.split("=", 1) for x in a.map)
        data = ingest(a.files, a.speaker, a.interviewer, names)
        Path(a.out).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        for f in data["files"]:
            note = f" · {f['no_timecode']} without timecode" if f["no_timecode"] else ""
            print(f"{f['letter']} {f['file']}: {f['bites']} bites · {f['duration']} · speakers: {', '.join(f['speakers'])}{note}")
        if any(b["speaker"] == "?" for b in data["bites"]):
            print("  warn  some lines have no speaker: rerun with --speaker NAME or fix the transcript")
        print(f"→ {a.out}")
        return
    data = json.loads(Path(a.bites).read_text(encoding="utf-8"))
    bites = data["bites"]
    if a.cmd == "index":
        for b in bites:
            if a.speaker and b["speaker"] != a.speaker:
                continue
            if a.grep and a.grep not in b["text"]:
                continue
            q = " (Q)" if b.get("interviewer") else ""
            text = b["text"] if len(b["text"]) <= a.width else b["text"][:a.width - 1] + "…"
            print(f"{b['id']} {b['tc_in']} {b['speaker']}{q}: {text}")
        return
    by = {b["id"]: b for b in bites}
    for i in expand(a.ids, bites):
        b = by.get(i)
        if not b:
            print(f"{i}: no such bite")
            continue
        d = f"  [{'; '.join(b['direction'])}]" if b.get("direction") else ""
        print(f"{b['id']} {b['file']} {b['tc_in']}–{b['tc_out']} {b['speaker']}: {b['text']}{d}")


if __name__ == "__main__":
    main()
