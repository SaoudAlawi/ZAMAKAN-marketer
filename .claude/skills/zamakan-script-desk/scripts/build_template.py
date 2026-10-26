#!/usr/bin/env python3
"""Build a Word template that helps a writer plan and write: foundations, characters, a beat sheet with page targets
for the chosen structure and running time, and a draft section set up with screenplay styles. No story text: only
questions and empty spaces to fill.

Usage:
  python build_template.py --structure three-act --minutes 90 --lang ar --out template.docx [--title "…"] [--pdf]
  python build_template.py --list                       # the structures and what they fit

Writing the draft in the template's styles (Scene Heading, Action, Character, Parenthetical, Dialogue, Transition)
lets format_script.py and the reader read it with no guessing.
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

import _doc as D  # noqa: E402
import _screenplay as S  # noqa: E402

T = {  # interface text
    "en": {"written_by": "Written by", "draft": "Draft", "date": "Date", "contact": "Contact", "how": "How to use this template",
           "how_items": ["Answer the questions in your own words; nothing here is written for you.",
                         "Page targets assume about one page per minute; treat them as a guide, not a rule.",
                         "Write the draft at the end using the styles (Scene Heading, Action, Character, Dialogue…), so it formats itself and can be read back for notes."],
           "foundations": "Foundations", "logline": "Logline", "logline_q": "In one or two sentences: who wants what, what stands in the way, and what is at stake?",
           "theme": "Theme", "theme_q": "What is this film about underneath the plot? Write it as a question, not a message.",
           "hero": "Protagonist", "hero_q": "Want (outside goal), need (inner lack), flaw, and what they do in the first scene.",
           "opp": "Opposition", "opp_q": "Who or what stands in the way, and what do they want?",
           "world": "World and tone", "world_q": "Where and when? What should the audience feel?",
           "characters": "Characters", "cols": ["Name", "Wants", "Needs", "Flaw", "How they change"],
           "beats": "Beat sheet", "timeline": "Timeline", "tcols": ["Beat", "Pages", "Share of the film"], "pages": "pages", "min": "min",
           "write": "Write here…", "draft_title": "Draft", "draft_hint": "Delete these sample lines and write your script in these styles.",
           "samples": [("scene_heading", "INT. PLACE - DAY"), ("action", "What we see and hear, in the present tense."),
                       ("character", "NAME"), ("parenthetical", "(how it is said, only if needed)"), ("dialogue", "What the character says."),
                       ("transition", "CUT TO:")], "source": "Source", "fits": "Suits", "kinds": {"feature": "feature films", "short": "short films", "documentary": "documentaries"},
           "proposal": "Proposal / treatment", "subjects": "Subjects", "subject_cols": ["Person", "Their part in the film", "Access", "What we film with them"],
           "av_title": "Two-column script (picture | sound)", "av_hint": "One row per moment: what we see, and what we hear beside it. Add rows as you go."},
    "ar": {"written_by": "تأليف", "draft": "المسودة", "date": "التاريخ", "contact": "للتواصل", "how": "كيف تستخدم هذا القالب",
           "how_items": ["أجب عن الأسئلة بكلماتك أنت؛ لا شيء هنا مكتوب نيابة عنك.",
                         "أهداف الصفحات مبنية على صفحة تقريباً لكل دقيقة؛ هي دليل لا قاعدة.",
                         "اكتب المسودة في آخر الملف باستخدام الأنماط (عنوان المشهد، الوصف، الشخصية، الحوار…) لتتنسق تلقائياً ويمكن قراءتها لاحقاً وإبداء الملاحظات."],
           "foundations": "الأساسات", "logline": "الفكرة في سطر", "logline_q": "في جملة أو جملتين: من يريد ماذا، وما الذي يقف في طريقه، وما الذي على المحك؟",
           "theme": "الموضوع", "theme_q": "عمّ يتحدث هذا الفيلم تحت سطح الحكاية؟ اكتبه سؤالاً لا رسالة.",
           "hero": "البطل", "hero_q": "رغبته (هدفه الظاهر)، وحاجته (نقصه الداخلي)، وعيبه، وماذا يفعل في المشهد الأول.",
           "opp": "القوة المضادة", "opp_q": "من أو ما الذي يقف في الطريق، وماذا يريد؟",
           "world": "العالم والنبرة", "world_q": "أين ومتى؟ وماذا تريد أن يشعر الجمهور؟",
           "characters": "الشخصيات", "cols": ["الاسم", "يريد", "يحتاج", "العيب", "كيف يتغيّر"],
           "beats": "محطات القصة", "timeline": "الخط الزمني", "tcols": ["المحطة", "الصفحات", "النسبة من الفيلم"], "pages": "صفحات", "min": "دقيقة",
           "write": "اكتب هنا…", "draft_title": "المسودة", "draft_hint": "احذف هذه الأسطر النموذجية واكتب نصك بهذه الأنماط.",
           "samples": [("scene_heading", "داخلي - المكان - نهار"), ("action", "ما نراه ونسمعه، بصيغة المضارع."),
                       ("character", "الاسم"), ("parenthetical", "(طريقة القول، عند الحاجة فقط)"), ("dialogue", "ما تقوله الشخصية."),
                       ("transition", "قطع إلى:")], "source": "المصدر", "fits": "يناسب", "kinds": {"feature": "الأفلام الطويلة", "short": "الأفلام القصيرة", "documentary": "الأفلام الوثائقية"},
           "proposal": "المقترح / المعالجة", "subjects": "الشخصيات (أشخاص حقيقيون)", "subject_cols": ["الشخص", "دوره في الفيلم", "الوصول إليه", "ما نصوّره معه"],
           "av_title": "سيناريو بعمودين (الصورة | الصوت)", "av_hint": "صف لكل لحظة: ما نراه وما نسمعه بجانبه. أضف صفوفاً كلما تقدمت."},
}
# a documentary is timed in screen minutes, not pages, and planned before (and again after) the shoot
DOC_T = {"en": {"written_by": "Director", "tcols": ["Beat", "Minutes", "Share of the film"], "pages": "min",
                "how_items": ["Answer the questions in your own words; nothing here is written for you.",
                              "Minute targets are a guide: a documentary is found in the shoot and shaped in the edit, so expect to redo this after filming.",
                              "Plan the film in the two-column table at the end (picture | sound); the paper edit and A/V tools read it back."]},
         "ar": {"written_by": "إخراج", "tcols": ["المحطة", "الدقائق", "النسبة من الفيلم"], "pages": "دقيقة",
                "how_items": ["أجب عن الأسئلة بكلماتك أنت؛ لا شيء هنا مكتوب نيابة عنك.",
                              "أهداف الدقائق دليل: الوثائقي يُكتشف في التصوير ويتشكّل في المونتاج، فتوقّع أن تعيد هذا بعد التصوير.",
                              "خطّط الفيلم في الجدول ذي العمودين في آخر الملف (الصورة | الصوت)؛ أدوات المونتاج الورقي والسيناريو بعمودين تقرؤه."]}}
GREY = RGBColor(0x8A, 0x8A, 0x8A)


def para(doc, text, lang, style=None, grey=False, size=None, bold=False, italic=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    if lang == "ar":
        r._r.get_or_add_rPr().append(S._el("w:rtl"))
    if grey:
        r.font.color.rgb = GREY
    if size:
        r.font.size = Pt(size)
        r._r.get_or_add_rPr().append(S._el("w:szCs", **{"w:val": int(size * 2) + 4}))
    r.bold, r.italic = bold or None, italic or None
    return p


def table(doc, rows, lang, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    pr = t._tbl.tblPr
    for old in pr.findall(S.qn("w:tblW")):
        pr.remove(old)
    pr.append(S._el("w:tblW", **{"w:w": 5000, "w:type": "pct"}))  # full text width
    if lang == "ar":
        pr.append(S._el("w:bidiVisual"))
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            p = cell.paragraphs[0]
            if lang == "ar":
                p._p.get_or_add_pPr().append(S._el("w:bidi"))
            r = p.add_run(val)
            if lang == "ar":
                r._r.get_or_add_rPr().append(S._el("w:rtl"))
            r.bold = (header and i == 0) or None
    return t


def write_box(doc, lang):
    t = table(doc, [[""]], lang, header=False)
    for _ in range(2):
        t.cell(0, 0).add_paragraph()
    p = t.cell(0, 0).paragraphs[0]
    r = p.add_run(T[lang]["write"])
    r.font.color.rgb = GREY
    if lang == "ar":
        r._r.get_or_add_rPr().append(S._el("w:rtl"))
    doc.add_paragraph()


def beat_pages(beat, minutes):
    """Pages a beat covers at about one page per minute: from the page it starts on to the page it ends on."""
    import math
    a = math.floor(beat["at"] * minutes / 100) + 1
    return min(a, minutes), max(min(a, minutes), math.ceil(beat["end"] * minutes / 100))


def build(structure_id, minutes, lang, title=None, paper=None):
    st = next((s for s in S.STRUCTURES["structures"] if s["id"] == structure_id), None)
    if not st:
        raise SystemExit(f"unknown structure {structure_id}; run --list")
    t, ar = T[lang], lang == "ar"
    name = st["name_ar"] if ar else st["name"]
    documentary = D.is_doc(structure_id)
    if documentary:
        t = {**t, **DOC_T[lang]}
    doc = S.new_document(lang, paper)
    # title page
    for _ in range(8):
        doc.add_paragraph()
    p = para(doc, title or ("عنوان الفيلم" if ar else "FILM TITLE"), lang, size=20, bold=True)
    p.alignment = 1
    for label in (t["written_by"], t["draft"], t["date"], t["contact"]):
        q = para(doc, f"{label}: ", lang)
        q.alignment = 1
    S.page_break(doc)
    S.add_page_numbers(doc, lang)
    # how to + structure summary
    para(doc, t["how"], lang, style="Heading 1")
    for item in t["how_items"]:
        para(doc, item, lang, style="List Bullet")
    para(doc, name, lang, style="Heading 2")
    para(doc, st["summary_ar"] if ar else st["summary"], lang)
    para(doc, f"{t['source']}: {st['source']} · {t['fits']}: {', '.join(t['kinds'][k] for k in st['fits'])} · {S.digits(minutes, lang)} {t['min']}", lang, grey=True, size=10)
    if documentary:  # the proposal / treatment sections funders ask for
        para(doc, t["proposal"], lang, style="Heading 1")
        for sec in D.PROPOSAL:
            para(doc, sec["name_ar"] if ar else sec["name"], lang, style="Heading 3")
            for q in (sec["q_ar"] if ar else sec["q"]):
                para(doc, q, lang, grey=True, italic=True)
            write_box(doc, lang)
        para(doc, t["subjects"], lang, style="Heading 1")
        table(doc, [t["subject_cols"]] + [[""] * 4 for _ in range(4)], lang)
    else:
        # foundations
        para(doc, t["foundations"], lang, style="Heading 1")
        for key in ("logline", "theme", "hero", "opp", "world"):
            para(doc, t[key], lang, style="Heading 3")
            para(doc, t[key + "_q"], lang, grey=True, italic=True)
            write_box(doc, lang)
        para(doc, t["characters"], lang, style="Heading 1")
        table(doc, [t["cols"]] + [[""] * 5 for _ in range(4)], lang)
    S.page_break(doc)
    # beat sheet: timeline, then one block per beat
    para(doc, f"{t['beats']}: {name}", lang, style="Heading 1")
    para(doc, t["timeline"], lang, style="Heading 3")
    rows = [t["tcols"]]
    for b in st["beats"]:
        a, z = beat_pages(b, minutes)
        rows.append([b["name_ar"] if ar else b["name"], f"{S.digits(a, lang)}–{S.digits(z, lang)}",
                     f"{S.digits(b['at'], lang)}–{S.digits(b['end'], lang)}%"])
    table(doc, rows, lang)
    doc.add_paragraph()
    for b in st["beats"]:
        a, z = beat_pages(b, minutes)
        para(doc, f"{b['name_ar'] if ar else b['name']}  ({t['pages']} {S.digits(a, lang)}–{S.digits(z, lang)})", lang, style="Heading 2")
        para(doc, b["purpose_ar"] if ar else b["purpose"], lang, italic=True)
        for q in (b["questions_ar"] if ar else b["questions"]):
            para(doc, q, lang, style="List Bullet")
        write_box(doc, lang)
    S.page_break(doc)
    if documentary:  # a two-column picture | sound draft
        para(doc, t["av_title"], lang, style="Heading 1")
        para(doc, t["av_hint"], lang, grey=True, italic=True)
        D.av_table(doc, [("", "", None)] * 8, lang)
        return doc
    # draft section in screenplay styles
    para(doc, t["draft_title"], lang, style="Heading 1")
    para(doc, t["draft_hint"], lang, grey=True, italic=True)
    for kind, text in t["samples"]:
        S.add_element(doc, kind, text, lang)
    return doc


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--structure", default="three-act")
    ap.add_argument("--minutes", type=int, default=None, help="planned running time (default 90 for features, 12 for shorts)")
    ap.add_argument("--lang", choices=["ar", "en"], default="ar")
    ap.add_argument("--paper", choices=["a4", "letter"])
    ap.add_argument("--title")
    ap.add_argument("--out")
    ap.add_argument("--pdf", action="store_true", help="also save a PDF preview")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--fits", choices=["short", "feature", "documentary"], help="with --list: only the structures for this kind of film")
    a = ap.parse_args()
    if a.list:
        for s in S.STRUCTURES["structures"]:
            if a.fits and a.fits not in s["fits"]:
                continue
            summary = s["summary_ar"] if a.lang == "ar" else s["summary"]  # one plain line to offer the writer
            print(f"{s['id']:18s} {s['name']} / {s['name_ar']} · fits {', '.join(s['fits'])} · beats: {', '.join(b['id'] for b in s['beats'])}\n{'':18s} {summary}")
        return
    st = next((s for s in S.STRUCTURES["structures"] if s["id"] == a.structure), None)
    minutes = a.minutes or (12 if st and st["fits"] == ["short"] else 25 if st and st["fits"] == ["documentary"] else 90)
    out = Path(a.out or f"template-{a.structure}-{a.lang}.docx")
    build(a.structure, minutes, a.lang, a.title, a.paper).save(out)
    print(f"template → {out} ({a.structure}, {minutes} min, {a.lang})")
    if a.pdf:
        pdf, err = S.to_pdf(out)
        print(f"preview → {pdf}" if pdf else err)


if __name__ == "__main__":
    main()
