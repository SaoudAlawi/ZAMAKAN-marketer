#!/usr/bin/env python3
"""Scene cards: the script as index cards, grouped by the beat each scene falls in, with the scene's mark (✕ ! ? ✓),
a one-line summary, who is in it, its length and what drives cost. And a way to try a new order: the writer
renumbers the cards in a Word table (or drops one with x), and the script is rebuilt in that order.

Usage:
  python cards.py board script.json --stats stats.json [--notes notes.json] --out <dir> [--no-pdf]   # HTML + A4 PDF board
  python cards.py table script.json [--notes notes.json] --out cards.docx                             # Word table to reorder
  python cards.py reorder script.json cards.docx --out script-reordered.json                          # read the new order
Then format_script.py / digest.py on script-reordered.json show the new shape (beat map, tension, cards).
"""
import argparse
import html
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _screenplay as S  # noqa: E402

E = html.escape
MARK = {"red": ("✕", "#d03b3b"), "amber": ("!", "#fab219"), "question": ("?", "#8a8984"), "green": ("✓", "#0ca30c")}
T = {"ar": {"title": "بطاقات المشاهد", "sub": "كل مشهد بطاقة، مجمّعة حسب المحطة التي يقع فيها. جرّب ترتيباً جديداً بجدول البطاقات في Word.",
            "scene": "مشهد", "order": "الترتيب الجديد", "head": "المشهد", "line": "في سطر", "who": "الشخصيات",
            "hint": "اكتب رقم الترتيب الجديد لكل مشهد (أو x لحذفه)، أو انقل الصفوف نفسها. ثم أرسل الملف.", "none": "خارج المحطات"},
     "en": {"title": "Scene cards", "sub": "Each scene a card, grouped by the beat it falls in. Try a new order with the cards table in Word.",
            "scene": "Scene", "order": "New order", "head": "Scene", "line": "In one line", "who": "Characters",
            "hint": "Write each scene's new position (or x to drop it), or move the rows themselves. Then send the file back.", "none": "outside the beats"}}


def one_liner(sc, n, notes):
    custom = ((notes or {}).get("cards") or {}).get(str(n))
    if custom:
        return custom
    first = next((e["text"] for e in sc["elements"] if e.get("scene") == n and e["type"] == "action"), "")
    return first[:90] + ("…" if len(first) > 90 else "")


def beat_of(stats):
    """Each scene's beat: the narrowest beat whose span holds the scene's middle."""
    out = {}
    beats = stats.get("beats") or []
    for r in stats["scene_rows"]:
        mid = r["pos"] + r["share"] / 2
        inside = [b for b in beats if b["at"] <= mid <= b["end"]]
        if inside:
            b = min(inside, key=lambda b: b["end"] - b["at"])
            out[r["n"]] = b["name"]
    return out


def board(sc, stats, notes, lang):
    from side_notes import scene_marks
    t = T[lang]
    marks = scene_marks(sc, notes) if notes else {}
    flags = {}
    for k, ns in (stats.get("production") or {}).items():
        for n in ns:
            flags.setdefault(n, []).append(k)
    from charts import FLAG, STRIP, strip_kind
    beats = beat_of(stats)
    groups = []
    for r in stats["scene_rows"]:
        g = beats.get(r["n"], t["none"])
        if not groups or groups[-1][0] != g:
            groups.append((g, []))
        groups[-1][1].append(r)
    title = (sc.get("title_page") or {}).get("title") or ""
    css = """:root{--bg:#f6f4ef;--ink:#1d1b18;--muted:#5e5a53;--card:#fff;--line:#e4dfd5}
@media (prefers-color-scheme:dark){:root:where(:not([data-theme=light])){--bg:#15140f;--ink:#f2efe8;--muted:#b5afa4;--card:#201f19;--line:#38352c}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.55 "IBM Plex Sans Arabic","Noto Naskh Arabic",system-ui,sans-serif}
.wrap{max-width:1180px;margin:0 auto;padding:28px 16px 56px}h1{font-size:28px;margin:0 0 4px}.sub{color:var(--muted);margin:0 0 18px}
.grp{margin:18px 0 6px;font-weight:700;font-size:15px;display:flex;align-items:center;gap:10px}.grp::after{content:"";flex:1;height:1px;background:var(--line)}
.cards{display:grid;grid-template-columns:repeat(auto-fill,minmax(210px,1fr));gap:12px}
.c{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:10px 12px;display:flex;flex-direction:column;gap:6px;break-inside:avoid;position:relative}
.c .top{display:flex;align-items:center;gap:8px;font-weight:700}.c .n{font-size:20px}.c .h{font-size:12.5px;color:var(--muted)}
.c .l{font-size:13.5px}.chips span{display:inline-block;border:1px solid var(--line);border-radius:999px;padding:0 8px;margin:0 0 3px 3px;font-size:11.5px}
.mk{display:inline-flex;width:20px;height:20px;border-radius:50%;align-items:center;justify-content:center;color:#fff;font:700 11px system-ui}.mk.amber{color:#1d1b18}
.strip{position:absolute;inset-inline-end:0;top:0;bottom:0;width:7px;border-radius:0 12px 12px 0;border-inline-start:1px solid var(--line)}
[dir=rtl] .strip{border-radius:12px 0 0 12px}.bar{height:5px;border-radius:3px;background:var(--muted);opacity:.4}.fl{font-size:11px;color:var(--muted)}
@page{size:A4 landscape;margin:10mm}@media print{body{background:#fff}.wrap{max-width:none;padding:0}.cards{grid-template-columns:repeat(5,1fr)}}"""
    mx = max((r["share"] for r in stats["scene_rows"]), default=1)
    h = [f'<!doctype html><html lang="{lang}" dir="{"rtl" if lang == "ar" else "ltr"}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
         f'<title>{E(title)} · {t["title"]}</title><link href="https://fonts.googleapis.com/css2?family=IBM+Plex+Sans+Arabic:wght@400;700&display=swap" rel="stylesheet">'
         f'<style>{css}</style></head><body><div class="wrap"><h1>{E(title)} · {t["title"]}</h1><p class="sub">{t["sub"]}</p>']
    for g, rs in groups:
        h.append(f'<div class="grp">{E(g)}</div><div class="cards">')
        for r in rs:
            m = marks.get(r["n"])
            mk = f'<b class="mk {m}" style="background:{MARK[m][1]}">{MARK[m][0]}</b>' if m else ""
            col, lab = STRIP[strip_kind(r)]
            fl = " · ".join(FLAG[lang].get(k, k) for k in flags.get(r["n"], []) if k != "night exterior")
            chips = "".join(f"<span>{E(c)}</span>" for c in r["speakers"])
            h.append(f'<div class="c"><i class="strip" style="background:{col}" title="{E(lab[lang])}"></i><div class="top"><span class="n">{S.digits(r["n"], lang)}</span>{mk}</div>'
                     f'<div class="h">{E(r.get("int_ext") or "")} - {E(r.get("place") or "")} - {E(r.get("time") or "")}</div>'
                     f'<div class="l">{E(one_liner(sc, r["n"], notes))}</div><div class="chips">{chips}</div>'
                     f'<div class="bar" style="width:{r["share"] / mx * 100:.0f}%"></div>' + (f'<div class="fl">{E(fl)}</div>' if fl else "") + "</div>")
        h.append("</div>")
    h.append("</div></body></html>")
    return "\n".join(h)


def table(sc, notes, lang):
    t = T[lang]
    doc = S.new_document(lang)
    p = doc.add_paragraph(style="Heading 1")
    r = p.add_run(((sc.get("title_page") or {}).get("title") or "") + " · " + t["title"])
    p = doc.add_paragraph()
    r = p.add_run(t["hint"])
    r.font.color.rgb = S.RGBColor(0x77, 0x77, 0x77)
    for x in (p, ):
        if lang == "ar":
            r._r.get_or_add_rPr().append(S._el("w:rtl"))
    who = {}
    for e in sc["elements"]:
        if e["type"] == "character" and e.get("scene"):
            who.setdefault(e["scene"], [])
            if e.get("name") not in who[e["scene"]]:
                who[e["scene"]].append(e.get("name"))
    tb = doc.add_table(rows=len(sc["scenes"]) + 1, cols=5)
    tb.style = "Table Grid"
    pr = tb._tbl.tblPr
    pr.append(S._el("w:tblW", **{"w:w": 5000, "w:type": "pct"}))
    pr.append(S._el("w:tblLayout", **{"w:type": "fixed"}))
    if lang == "ar":
        pr.append(S._el("w:bidiVisual"))
    for gc, w in zip(tb._tbl.tblGrid.findall(S.qn("w:gridCol")), (0.45, 0.8, 1.6, 2.2, 1.0)):
        gc.set(S.qn("w:w"), str(int(w * 1440)))
    heads = ["#", t["order"], t["head"], t["line"], t["who"]]
    rows = [heads] + [[str(s["n"]), "", f'{s.get("int_ext") or ""} - {s.get("place") or ""} - {s.get("time") or ""}',
                       one_liner(sc, s["n"], notes), "، ".join(who.get(s["n"], [])) if lang == "ar" else ", ".join(who.get(s["n"], []))] for s in sc["scenes"]]
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cp = tb.cell(i, j).paragraphs[0]
            if lang == "ar":
                cp._p.get_or_add_pPr().append(S._el("w:bidi"))
            rr = cp.add_run(val)
            rr.bold = (i == 0 or j == 0) or None
            if lang == "ar":
                rr._r.get_or_add_rPr().append(S._el("w:rtl"))
    return doc


def read_order(path, nscenes):
    """New order from the cards table: numbers in the 'new order' column win; else the row order; 'x' drops a scene."""
    from docx import Document
    tb = Document(path).tables[0]
    rows = []
    for k, row in enumerate(tb.rows[1:]):
        cells = [c.text.strip() for c in row.cells]
        num = cells[0].translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))
        if not num.isdigit():
            continue
        new = cells[1].translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")).strip().lower()
        if new in ("x", "×", "حذف"):
            continue
        rows.append((float(new) if re.fullmatch(r"\d+(\.\d+)?", new) else k + 1 + 0.001 * k, k, int(num)))
    order = [n for _, _, n in sorted(rows)]
    dropped = sorted(set(range(1, nscenes + 1)) - set(order))
    return order, dropped


def reorder(sc, order):
    """A copy of the script with its scenes in the new order (title page and anything before scene 1 kept)."""
    pre = [e for e in sc["elements"] if not e.get("scene")]
    by = {}
    for e in sc["elements"]:
        if e.get("scene"):
            by.setdefault(e["scene"], []).append(dict(e))
    els = pre[:]
    for new_n, old_n in enumerate(order, 1):
        for e in by.get(old_n, []):
            e["scene"] = new_n
            e["was_scene"] = old_n
            els.append(e)
    import parse_script as P
    out = {k: v for k, v in sc.items() if k not in ("elements", "scenes", "format")}
    out["elements"] = els
    out["scenes"] = P.number_scenes(els, sc["lang"])
    for s in out["scenes"]:
        s["was"] = order[s["n"] - 1]
    out["reordered_from"] = order
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("board")
    b.add_argument("script")
    b.add_argument("--stats", required=True)
    b.add_argument("--notes")
    b.add_argument("--out", default=".")
    b.add_argument("--no-pdf", action="store_true")
    tb = sub.add_parser("table")
    tb.add_argument("script")
    tb.add_argument("--notes")
    tb.add_argument("--out", required=True)
    r = sub.add_parser("reorder")
    r.add_argument("script")
    r.add_argument("cards")
    r.add_argument("--out", required=True)
    a = ap.parse_args()
    load = lambda p: json.loads(Path(p).read_text(encoding="utf-8")) if p else None  # noqa: E731
    sc = load(a.script)
    lang = sc["lang"]
    stem = re.sub(r"[^\w؀-ۿ-]+", "-", (sc.get("title_page") or {}).get("title") or "script").strip("-")
    if a.cmd == "board":
        out = Path(a.out)
        out.mkdir(parents=True, exist_ok=True)
        path = out / f"{stem}-cards.html"
        path.write_text(board(sc, load(a.stats), load(a.notes), lang), encoding="utf-8")
        print(f"cards → {path}")
        if not a.no_pdf:
            from side_notes import to_pdf
            print(f"PDF → {to_pdf(path)}")
    elif a.cmd == "table":
        table(sc, load(a.notes), lang).save(a.out)
        print(f"cards table → {a.out} (the writer fills 'new order' or moves rows; read it back with: cards.py reorder)")
    else:
        order, dropped = read_order(a.cards, len(sc["scenes"]))
        new = reorder(sc, order)
        Path(a.out).write_text(json.dumps(new, ensure_ascii=False, indent=1), encoding="utf-8")
        moved = [f"{old}→{i}" for i, old in enumerate(order, 1) if old != i]
        print(f"new order: {' '.join(map(str, order))}" + (f" · dropped {dropped}" if dropped else "") + f"\nmoved: {', '.join(moved) or 'none'} → {a.out}")


if __name__ == "__main__":
    main()
