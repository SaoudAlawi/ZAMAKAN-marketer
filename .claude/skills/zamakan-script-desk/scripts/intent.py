#!/usr/bin/env python3
"""The writer's intent sheet: what the film is about, what they are unsure of, which choices are deliberate, each
character's dialect, where it should play. Notes are then judged against what the writer is going for (the way good
script editors read), and declared choices (a character who switches to formal Arabic on purpose) are not flagged
as mistakes.

Usage:
  python intent.py form --lang ar --title "المرساة" [--script script.json] --out intent.docx   # a Word form to fill
  python intent.py read intent.docx --out intent.json                                          # read the filled form
Claude can also write intent.json straight from the conversation (same keys; see FIELDS).
"""
import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _screenplay as S  # noqa: E402

FIELDS = [  # key, Arabic label, English label, hint (ar), hint (en)
    ("about", "عمّ يتحدث فيلمك؟", "What is your film about?", "في سطر: الحكاية والمعنى تحتها", "One line: the story and what lies under it"),
    ("question", "ما السؤال الذي يطرحه الفيلم؟", "What question does the film ask?", "سؤال لا رسالة", "A question, not a message"),
    ("feel", "ماذا تريد أن يشعر الجمهور عند النهاية؟", "What should the audience feel at the end?", "", ""),
    ("unsure", "ما الذي لست متأكداً منه؟", "What are you unsure of?", "مشاهد أو قرارات تريد رأياً فيها", "Scenes or choices you want a view on"),
    ("intentional", "اختيارات مقصودة لا تريدها أن تُعدّ أخطاء", "Deliberate choices (don't flag these as mistakes)", "مثل: الجدة تتكلم بالفصحى حين تبارك؛ النهاية مفتوحة عمداً", "e.g. the grandmother speaks formal Arabic when she blesses; the ending is open on purpose"),
    ("audience", "لمن الفيلم وأين يُعرض؟", "Who is it for, and where should it play?", "مهرجانات، منصات، تلفزيون، سوق", "Festivals, platforms, TV, market"),
    ("minutes", "المدة المخططة بالدقائق", "Planned running time (minutes)", "", ""),
    ("notes_wanted", "أي نوع من الملاحظات تريد الآن؟", "What kind of notes do you want now?", "قراءة كاملة / الحوار فقط / أول ١٠ صفحات / زاوية واحدة", "Full read / dialogue only / first 10 pages / one viewpoint"),
    ("references", "أفلام يتحاور معها فيلمك", "Films yours is in conversation with", "اختياري", "Optional"),
]
CHAR_COLS = {"ar": ["الشخصية", "اللهجة", "ماذا تريد", "اختيارات مقصودة في كلامها أو سلوكها"],
             "en": ["Character", "Dialect", "What they want", "Deliberate choices in how they speak or act"]}
DIALECT_WORDS = [("mixed", r"يتنقل|تتنقل|مختلط|mix|switch"), ("msa", r"فصح|msa|standard|formal"), ("gulf", r"خليج|عمان|عُمان|gulf|omani|khaleeji"),
                 ("egyptian", r"مصر|egypt"), ("levantine", r"شام|لبنان|سوري|levant")]


def dialect_code(text):
    for code, rx in DIALECT_WORDS:
        if re.search(rx, text or "", re.I):
            return code
    return None


def run(p, text, lang, bold=False, grey=False, size=None):
    r = p.add_run(text)
    r.bold = bold or None
    if grey:
        r.font.color.rgb = S.RGBColor(0x8A, 0x8A, 0x8A)
    if size:
        r.font.size = S.Pt(size)
    if lang == "ar":
        r._r.get_or_add_rPr().append(S._el("w:rtl"))
    return r


def cell_par(cell, lang):
    p = cell.paragraphs[0]
    if lang == "ar":
        p._p.get_or_add_pPr().append(S._el("w:bidi"))
    return p


def table(doc, rows, cols, lang, widths):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    pr = t._tbl.tblPr
    pr.append(S._el("w:tblW", **{"w:w": 5000, "w:type": "pct"}))
    pr.append(S._el("w:tblLayout", **{"w:type": "fixed"}))
    if lang == "ar":
        pr.append(S._el("w:bidiVisual"))
    for gc, w in zip(t._tbl.tblGrid.findall(S.qn("w:gridCol")), widths):
        gc.set(S.qn("w:w"), str(int(w * 1440)))
    return t


def form(lang, title, script=None):
    ar = lang == "ar"
    doc = S.new_document(lang)
    p = doc.add_paragraph(style="Heading 1")
    run(p, ("ورقة النية: " if ar else "Intent sheet: ") + (title or ""), lang)
    p = doc.add_paragraph()
    run(p, "اكتب بكلماتك. هذه الورقة تساعد القارئ على الحكم على نصك بما تريده أنت، لا بذوقه. اترك ما لا تعرفه فارغاً." if ar else
        "Write in your own words. This sheet lets the reader judge your script by what you are going for, not by their taste. Leave blank what you don't know.", lang, grey=True, size=10)
    t = table(doc, len(FIELDS), 2, lang, (2.1, 3.7))
    for i, (key, la, le, ha, he) in enumerate(FIELDS):
        lp = cell_par(t.cell(i, 0), lang)
        run(lp, la if ar else le, lang, bold=True)
        hint = ha if ar else he
        if hint:
            run(t.cell(i, 0).add_paragraph(), hint, lang, grey=True, size=9)
        cell_par(t.cell(i, 1), lang)
        trpr = t.rows[i]._tr.get_or_add_trPr()
        trpr.append(S._el("w:trHeight", **{"w:val": 620, "w:hRule": "atLeast"}))  # room to write
    doc.add_paragraph()
    p = doc.add_paragraph(style="Heading 2")
    run(p, "الشخصيات" if ar else "Characters", lang)
    names = []
    if script:
        from collections import Counter
        names = [n for n, _ in Counter(e.get("name") for e in script["elements"] if e["type"] == "character").most_common(8)]
    rows = max(len(names), 4) + 1
    ct = table(doc, rows, 4, lang, (1.2, 1.1, 1.6, 1.9))
    for j, h in enumerate(CHAR_COLS[lang]):
        run(cell_par(ct.cell(0, j), lang), h, lang, bold=True)
    for i in range(1, rows):
        for j in range(4):
            cp = cell_par(ct.cell(i, j), lang)
            if j == 0 and i - 1 < len(names):
                run(cp, names[i - 1], lang)
    p = doc.add_paragraph()
    run(p, "اللهجة: خليجي/عماني، فصحى، مصري، شامي، أو «يتنقل» إن كانت تتنقل عمداً." if ar else
        "Dialect: Gulf/Omani, MSA, Egyptian, Levantine, or 'switches' if they switch on purpose.", lang, grey=True, size=9)
    return doc


def read(path):
    from docx import Document
    d = Document(path)
    if len(d.tables) < 2:
        raise SystemExit("not an intent sheet (expected the answers table and the characters table)")
    out = {"schema": "zamakan-script-intent/1"}
    t = d.tables[0]
    for i, (key, *_rest) in enumerate(FIELDS):
        if i < len(t.rows):
            val = "\n".join(p.text.strip() for p in t.cell(i, 1).paragraphs if p.text.strip())
            if val:
                out[key] = val
    for key in ("unsure", "intentional"):
        if key in out:
            out[key] = [x.strip(" -•·") for x in re.split(r"\n|؛|;", out[key]) if x.strip(" -•·")]
    if "minutes" in out:
        m = re.search(r"[\d٠-٩]+", out["minutes"])
        out["minutes"] = int(m.group(0).translate(str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789"))) if m else None
    chars = []
    for row in d.tables[1].rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        if cells[0]:
            chars.append({"name": cells[0], "dialect": dialect_code(cells[1]), "dialect_text": cells[1] or None,
                          "wants": cells[2] or None, "deliberate": cells[3] or None})
    if chars:
        out["characters"] = chars
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    f = sub.add_parser("form")
    f.add_argument("--lang", choices=["ar", "en"], default="ar")
    f.add_argument("--title")
    f.add_argument("--script")
    f.add_argument("--out", required=True)
    r = sub.add_parser("read")
    r.add_argument("docx")
    r.add_argument("--out", required=True)
    a = ap.parse_args()
    if a.cmd == "form":
        sc = json.loads(Path(a.script).read_text(encoding="utf-8")) if a.script else None
        form(a.lang, a.title, sc).save(a.out)
        print(f"intent sheet → {a.out} (send it to the writer; read it back with: intent.py read {a.out} --out intent.json)")
    else:
        data = read(a.docx)
        Path(a.out).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        filled = [k for k, *_ in FIELDS if data.get(k)]
        print(f"intent → {a.out}: {len(filled)}/{len(FIELDS)} answers, {len(data.get('characters', []))} characters"
              + "".join(f"\n  {c['name']}: {c['dialect'] or '?'}" + (f" · deliberate: {c['deliberate']}" if c.get("deliberate") else "") for c in data.get("characters", [])))


if __name__ == "__main__":
    main()
