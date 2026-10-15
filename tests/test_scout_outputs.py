"""Scout output builders: territories.md, report.md, board.html, deck.pptx from references.json."""
import html
import json
import math
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / ".claude/skills/zamakan-reference-scout/scripts"
FIX = REPO / "evals/fixtures/scout"
sys.path.insert(0, str(SCRIPTS))

import build_outputs as B  # noqa: E402


@pytest.fixture
def run_path(tmp_path):
    run = json.loads((FIX / "rawaq-ramadan-sample-run.json").read_text(encoding="utf-8"))
    for r in run["references"]:
        r.pop("thumbnail", None)  # offline + fast: no image downloads in unit tests
    run["references"][3]["status"] = "rejected"  # SAIB greeting: must never be shown
    p = tmp_path / "references.json"
    p.write_text(json.dumps(run, ensure_ascii=False), encoding="utf-8")
    return p


def build(run_path, tmp_path, formats="board,deck,report"):
    out = tmp_path / "job"
    sys.argv = ["build_outputs.py", str(run_path), "--out", str(out), "--formats", formats,
                "--brief", str(FIX / "rawaq-ramadan-brief.json"), "--client-name", "Rawaq Coffee Roasters", "--no-pdf"]
    B.main()
    return out, json.loads(run_path.read_text(encoding="utf-8"))


def test_territories_md_always_built(run_path, tmp_path):
    out, run = build(run_path, tmp_path, formats="")
    t = (out / "territories.md").read_text(encoding="utf-8")
    for terr in run["territories"]:
        assert terr["name"] in t and terr["name_ar"] in t
    assert "## Recommendation" in t and "- R0" in t
    assert "https://www.adsoftheworld.com/campaigns/iftar-is-always-better-together" in t
    assert not (out / "board.html").exists()


def test_report_sections_and_no_rejected(run_path, tmp_path):
    out, run = build(run_path, tmp_path)
    r = (out / "report.md").read_text(encoding="utf-8")
    for h in ("## Summary", "## Patterns", "## White space", "## Creative territories", "## Deep dives", "## All references"):
        assert h in r
    assert "SAIB" not in r
    n = sum(x["status"] != "rejected" for x in run["references"])
    assert f"{n} references ({n} verified" in r


def test_board_cards_filters_and_rtl(run_path, tmp_path):
    out, run = build(run_path, tmp_path)
    h = (out / "board.html").read_text(encoding="utf-8")
    shown = [x for x in run["references"] if x["status"] != "rejected"]
    assert h.count('class="card"') == len(shown) == len(run["references"]) - 1
    assert "ramadan-amp-eid-greeting-2024-saib-bank" not in h
    for ring in ("direct", "category", "adjacent", "wildcard"):
        assert f'data-ring="{ring}"' in h and f'data-v="{ring}"' in h
    assert 'class="ar-d" lang="ar"' in h and "مين يعزم الليلة؟" in h
    assert h.startswith('<meta charset="utf-8">') and "--bg:#F6F6F2" in h  # house style when the client has no design system
    assert h.count('target="_blank" rel="noopener"') >= len(shown)
    # strategy layer, in kickoff order
    order = [h.index(f'id="{s}"') for s in ("insights", "recommendation", "territories", "sparks", "competitors", "patterns", "references", "kickoff")]
    assert order == sorted(order)
    for x in run["insights"]:
        assert html.escape(x["source_url"]) in h
    for c in run["competitor_summary"]:
        assert html.escape(c["advertiser"]) in h
    codes = B.ref_codes({"references": shown})
    assert set(codes.values()) == {f"R{i:02d}" for i in range(1, len(shown) + 1)}
    assert all(f'<span class="code">{c}</span>' in h for c in codes.values())


def test_deck_structure_and_links(run_path, tmp_path):
    from pptx import Presentation
    out, run = build(run_path, tmp_path)
    prs = Presentation(out / "deck.pptx")
    n_refs = sum(r["status"] != "rejected" for r in run["references"])
    # cover, at a glance, insights, recommendation, territories, competitors, patterns, deep dives, refs (8/slide), kickoff, close
    expected = 8 + len(run["territories"]) + len(run["top"]) + math.ceil(n_refs / 8) + any(r.get("spark") for r in run["references"])
    assert len(prs.slides) == expected
    links = [r.hyperlink.address for s in prs.slides for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    assert len(links) >= n_refs
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Who hosts tonight?" in text and "SAIB" not in text
    assert "SPARKS FROM ELSEWHERE" in text and "Insights" in text and "What competitors run" in text and "Before kickoff" in text and "R01" in text


def test_refuses_invalid_run(tmp_path):
    run = json.loads((FIX / "rawaq-ramadan-sample-run.json").read_text(encoding="utf-8"))
    run["territories"][0]["reference_ids"].append("invented-campaign")
    p = tmp_path / "bad.json"
    p.write_text(json.dumps(run))
    with pytest.raises(SystemExit, match="unknown reference id 'invented-campaign'"):
        build(p, tmp_path)


def test_display_title_fallbacks():
    run = {"run": {"id": "x", "question": "The brief is in brief.md. Find us references…", "client": "rawaq"}}
    assert B.display_title(run, {"moment": "Ramadan 1448 (≈8 Feb–9 Mar 2027)"}, "Rawaq Coffee Roasters") == \
        "Rawaq Coffee Roasters · Ramadan 1448 references"
    run["run"]["title"] = "Rawaq · Ramadan 2027 references"
    assert B.display_title(run, None, None) == "Rawaq · Ramadan 2027 references"
