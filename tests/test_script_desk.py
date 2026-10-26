"""Script desk: templates, parser, formatter, digest and the reader's outputs (no network; LibreOffice tests skip
when it is missing)."""
import json
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
DESK = ROOT / ".claude" / "skills" / "zamakan-script-desk"
sys.path.insert(0, str(DESK / "scripts"))
import _screenplay as S  # noqa: E402
import annotate as A  # noqa: E402
import build_template as B  # noqa: E402
import digest as D  # noqa: E402
import format_script as F  # noqa: E402
import parse_script as P  # noqa: E402

SAMPLE = ROOT / "samples" / "script-desk" / "al-mirsah-draft.txt"
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


@pytest.fixture(scope="module")
def sc():
    s = P.parse(SAMPLE, fixes={"56": "character"})
    return s


def test_structures_are_well_formed():
    for st in S.STRUCTURES["structures"]:
        assert st["name_ar"] and st["source"] and st["fits"]
        for b in st["beats"]:
            assert 0 <= b["at"] < b["end"] <= 100, (st["id"], b["id"])
            assert b["questions"] and len(b["questions"]) == len(b["questions_ar"]), (st["id"], b["id"])


def test_beat_pages_cover_the_film():
    st = next(s for s in S.STRUCTURES["structures"] if s["id"] == "short-one-moment")
    pages = [B.beat_pages(b, 12) for b in st["beats"]]
    assert pages[0][0] == 1 and pages[-1][1] == 12
    assert all(a <= b for a, b in pages)


def test_template_has_styles_and_rtl(tmp_path):
    doc = B.build("three-act", 90, "ar", "اختبار")
    names = {s.name for s in doc.styles}
    assert {"Scene Heading", "Action", "Character", "Dialogue", "Parenthetical", "Transition"} <= names
    assert doc.sections[0]._sectPr.find(S.qn("w:bidi")) is not None
    assert doc.sections[0].right_margin > doc.sections[0].left_margin  # binding margin on the right in Arabic
    out = tmp_path / "t.docx"
    doc.save(out)
    back = P.parse(out)  # the template's sample draft lines read back by style
    assert len(back["scenes"]) == 1 and any(e["type"] == "character" for e in back["elements"])


def test_parse_arabic_sample(sc):
    assert sc["lang"] == "ar" and len(sc["scenes"]) == 13
    assert sc["title_page"]["title"] == "المرساة"
    s4 = [e for e in sc["elements"] if e.get("scene") == 4]
    assert [e["type"] for e in s4[:4]] == ["scene_heading", "action", "character", "dialogue"]  # «الجدة: …» split
    assert sc["scenes"][6]["int_ext"] == "خارجي" and sc["scenes"][6]["time"] == "ليل"
    assert not sc["unsure"]


def test_unsure_one_off_cue_is_listed_not_guessed():
    s = P.parse(SAMPLE)
    assert [u["text"] for u in s["unsure"]] == ["صياد"]


def test_parse_english_fountain(tmp_path):
    f = tmp_path / "x.fountain"
    f.write_text("INT. KITCHEN - NIGHT\n\nA kettle.\n\nSARA\n(quietly)\nNot tonight.\n\nJOHN (O.S.)\nSara?\n\nCUT TO:\n\n.ROOFTOP\n\nRain.\n", encoding="utf-8")
    s = P.parse(f)
    types = [e["type"] for e in s["elements"]]
    assert s["lang"] == "en" and len(s["scenes"]) == 2
    assert types.count("character") == 2 and "parenthetical" in types and "transition" in types
    john = next(e for e in s["elements"] if e.get("name") == "JOHN")
    assert john["ext"] == "O.S."


def test_heading_parts_any_order():
    assert P.heading_parts("مشهد ٣: ليل - داخلي - بيت سالم") == {"int_ext": "داخلي", "place": "بيت سالم", "time": "ليل"}
    assert P.heading_parts("EXT. STREET - DAY")["int_ext"].startswith("EXT")


def test_format_round_trip(sc, tmp_path):
    doc = F.build_docx(sc)
    out = tmp_path / "f.docx"
    doc.save(out)
    back = P.parse(out)
    assert len(back["scenes"]) == 13 and back["title_page"]["title"] == "المرساة"
    assert [e["type"] for e in back["elements"]] == [e["type"] for e in sc["elements"] if e["type"] not in ("note", "synopsis", "section")]
    assert "@صياد" in F.to_fountain(sc) and ".خارجي - شاطئ صور - فجر #1#" in F.to_fountain(sc)


@pytest.mark.skipif(not SOFFICE, reason="LibreOffice not installed")
def test_pdf_embeds_fonts_and_mirrors(sc, tmp_path):
    import pymupdf
    out = tmp_path / "f.docx"
    F.build_docx(sc).save(out)
    pdf, err = S.to_pdf(out, tmp_path)
    assert pdf, err
    with pymupdf.open(pdf) as d:
        fonts = {f[3] for p in d for f in p.get_fonts()}
        assert any("Amiri" in f for f in fonts)
        page = d[1]
        x1 = max(l["bbox"][2] for b in page.get_text("dict")["blocks"] for l in b.get("lines", []))
        assert abs(page.rect.width - x1 - 1.5 * 72) < 12  # text ends at the 1.5" binding margin on the right


def test_digest_flags(sc):
    d = D.digest(sc, "short-one-moment")
    assert d["protagonist_guess"] == "مريم"
    assert [3, 5] in d["same_place_same_people"]
    assert any(m["character"] == "الجدة" for m in d["dialect_mix"])
    assert any(t["scene"] == 1 for t in d["telling_lines"])
    assert 7 in d["production"]["night exterior"] and 6 in d["silent_scenes"]
    assert len(D.show(d)) < 6000  # stays a compact read


def notes(**over):
    n = {"schema": "zamakan-script-notes/1", "lang": "ar", "logline_read": "x", "overall": {"verdict": "amber", "summary": "x"},
         "lenses": [{"id": lid, "verdict": "green", "summary": "x"} for lid in A.LENS],
         "notes": [{"scene": 1, "lens": "visual", "kind": "strength", "text": "صورة افتتاحية قوية."}], "next_draft": ["x"]}
    n.update(over)
    return n


def test_notes_checker(sc):
    errs, warns = A.check(sc, notes())
    assert errs == [] and any("tension" in w for w in warns)
    bad = notes()
    bad["lenses"][1]["verdict"] = "red"  # red with no note
    bad["notes"].append({"scene": 99, "lens": "structure", "kind": "issue", "text": "x"})
    bad["notes"].append({"scene": 2, "lens": "dialogue", "kind": "issue", "text": "x", "quote": "كلام غير موجود"})
    errs, _ = A.check(sc, bad)
    assert any("scene 99" in e for e in errs) and any("quote not found" in e for e in errs)
    missing = notes(lenses=[{"id": "premise", "verdict": "green", "summary": "x"}])
    assert any("has no verdict" in e for e in A.check(sc, missing)[0])
    rewrite = notes(notes=[{"scene": 2, "lens": "dialogue", "kind": "issue", "text": "اكتبها هكذا: «أنا خايفة عليك يا بنتي من البحر ومن كل شي فيه وما أقدر أشوفك تروحين»"}])
    assert any("rewriting" in w for w in A.check(sc, rewrite)[1])


def test_annotated_docx_has_comments_and_notes_page(sc, tmp_path):
    n = notes()
    n["lenses"][1]["verdict"] = "amber"
    n["notes"].append({"scene": 3, "lens": "structure", "kind": "question", "text": "ما الجديد هنا؟", "quote": "خلاص. قلت لا."})
    doc = A.annotate(sc, n, "ar")
    out = tmp_path / "n.docx"
    doc.save(out)
    from docx import Document
    d = Document(out)
    assert len(d.comments) == 2
    texts = [p.text for p in d.paragraphs]
    assert texts.index("ملاحظات القراءة") < next(i for i, t in enumerate(texts) if "خارجي - شاطئ صور" in t)


def test_cli_end_to_end(tmp_path):
    js = tmp_path / "s.json"
    r = subprocess.run([sys.executable, str(DESK / "scripts" / "parse_script.py"), str(SAMPLE), "--out", str(js)], capture_output=True, text=True)
    assert r.returncode == 0 and "unsure" in r.stdout
    r = subprocess.run([sys.executable, str(DESK / "scripts" / "show.py"), str(js), "9"], capture_output=True, text=True)
    assert "الجدة:" in r.stdout and "S9" in r.stdout


# ---- v0.2: charts, side notes, voices, draft comparison ----
import charts as C  # noqa: E402
import compare as CMP  # noqa: E402
import side_notes as SN  # noqa: E402

SAMPLE2 = ROOT / "samples" / "script-desk" / "al-mirsah-draft2.txt"


def full_notes():
    n = notes()
    n["structure"] = "short-one-moment"
    n["lenses"][2]["verdict"] = "red"  # character
    n["notes"] += [{"scene": 8, "lens": "character", "kind": "issue", "text": "القرار قرار سالم. ماذا تختار مريم؟", "quote": "مريم بتطلع مع حمد"},
                   {"scene": 4, "lens": "structure", "kind": "question", "text": "ما دور البوصلة؟"}]
    n["lenses"][1]["verdict"] = "amber"
    n["turns"] = {"hook": 1, "turn": 8, "after": 13}
    n["tension"] = {str(i): (5 if i == 7 else 2) for i in range(1, 14)}
    return n


def test_turns_and_tension_are_checked(sc):
    n = full_notes()
    assert A.check(sc, n)[0] == []
    n["turns"]["midpoint"] = 5  # not a beat of short-one-moment
    n["tension"]["40"] = 3
    errs = A.check(sc, n)[0]
    assert any("not a beat" in e for e in errs) and any("unknown scenes" in e for e in errs)


def test_scene_marks_worst_wins(sc):
    m = SN.scene_marks(sc, full_notes())
    assert m[8] == "red" and m[4] == "question" and m[1] == "green"


def test_side_notes_place_cards_beside_their_lines(sc):
    page = SN.build(sc, full_notes(), "ar")
    assert page.count('class="card') == 3 and 'dir="rtl"' in page
    i = page.index("مريم بتطلع مع حمد")  # the quoted line, then its card in the same row
    assert page.index("القرار قرار سالم", i) - i < 400


def test_charts_render_for_arabic(sc):
    d = D.digest(sc, "short-one-moment")
    st = next(s for s in S.STRUCTURES["structures"] if s["id"] == "short-one-moment")
    frags = C.all_charts(d, full_notes(), st, SN.scene_marks(sc, full_notes()), "ar")
    assert len(frags) == 6
    html = "".join(frags)
    assert "خريطة المحطات" in html and "منحنى التوتر" in html and "خريطة اللهجات" in html and "لوحة التصوير" in html
    assert "#d03b3b" in html  # the red scene in the health strip
    assert C.strip_kind({"int_ext": "خارجي", "time": "ليل"}) == "ext_night" and C.strip_kind({"int_ext": "داخلي", "time": "فجر"}) == "dawn_dusk"


def test_voice_profiles(sc):
    d = D.digest(sc)
    v = {x["name"]: x for x in d["voices"]}
    assert v["الجدة"]["avg_words"] > v["مريم"]["avg_words"] * 2
    assert all("،" not in w and "؟" not in w for x in d["voices"] for w in x["own_words"])


def test_compare_drafts(sc):
    sc2 = P.parse(SAMPLE2)
    c = CMP.compare(sc, full_notes(), sc2)
    assert c["cut"] == [5, 7, 9] and c["scene_map"]["1"]["status"] == "same"
    assert 1 not in c["reread"] and 7 in c["reread"]
    by = {it["old_scene"]: it for it in c["previous_notes"]}
    assert by[8]["new_scene"] == 7 and by[8]["start"].startswith("re-read")
    assert by[4]["start"].startswith("probably open")  # compass scene unchanged: Claude must still check elsewhere


def test_previous_section_in_outputs(sc):
    n = full_notes()
    n["previous"] = {"draft": "1", "items": [{"id": 0, "status": "resolved", "text": "قرار مريم", "scene": 7},
                                             {"id": 1, "status": "open", "text": "سطر الأفكار في المشهد ١", "scene": 1}]}
    assert A.check(sc, n)[0] == []
    import report as R
    page = R.build(sc, n, D.digest(sc, "short-one-moment"), "ar")
    assert "منذ المسودة 1" in page and "حُلّت" in page and "ما زالت قائمة" in page


# ---- v0.3: intent sheet, margin conversation, scene cards, arcs ----
import cards as CARDS  # noqa: E402
import intent as INT  # noqa: E402
import replies as REP  # noqa: E402


def filled_intent(tmp_path, sc):
    from docx import Document
    path = tmp_path / "intent.docx"
    INT.form("ar", "المرساة", sc).save(path)
    d = Document(path)
    t = d.tables[0]
    t.cell(0, 1).paragraphs[0].add_run("عن الحزن حين يتحول إلى فعل")
    t.cell(4, 1).paragraphs[0].add_run("الجدة تبارك بالفصحى عمداً")
    t.cell(4, 1).add_paragraph("النهاية مفتوحة عمداً")
    t.cell(6, 1).paragraphs[0].add_run("١٢ دقيقة")
    for row in d.tables[1].rows[1:]:
        name = row.cells[0].text.strip()
        row.cells[1].paragraphs[0].add_run({"الجدة": "يتنقل", "الأم": "خليجي"}.get(name, ""))
    d.save(path)
    return INT.read(path)


def test_intent_round_trip_and_declared_dialects(tmp_path, sc):
    it = filled_intent(tmp_path, sc)
    assert it["about"].startswith("عن الحزن") and it["intentional"] == ["الجدة تبارك بالفصحى عمداً", "النهاية مفتوحة عمداً"]
    assert it["minutes"] == 12 and {c["name"]: c["dialect"] for c in it["characters"]}["الجدة"] == "mixed"
    d = D.digest(sc, intent=it)
    flagged = {m["character"]: m for m in d["dialect_mix"]}
    assert "الجدة" not in flagged  # declared as switching on purpose
    assert flagged["الأم"]["declared"] == "gulf" and "2" in flagged["الأم"]["by_scene"]
    assert "writer's intent" in D.show(d)


def test_intent_check_required_when_intent_given(sc):
    n = full_notes()
    assert any("intent_check" in e for e in A.check(sc, n, intent={"about": "x"})[0])
    n["intent_check"] = [{"intent": "x", "lands": "partly", "note": "y", "scene": 8}]
    assert A.check(sc, n, intent={"about": "x"})[0] == []


def test_replies_read_back_from_word(tmp_path, sc):
    n = full_notes()
    doc = A.annotate(sc, n, "ar")
    target = next(p for p in doc.paragraphs if "مريم بتطلع مع حمد" in p.text)
    doc.add_comment(target.runs, text="معك حق، بغيّرها.", author="الكاتب", initials="K")
    out = tmp_path / "back.docx"
    doc.save(out)
    items = REP.read(out, n)
    by = {i["id"]: i for i in items}
    k = next(i for i, x in enumerate(n["notes"]) if x.get("quote") == "مريم بتطلع مع حمد")
    assert by[k]["replies"][0]["text"] == "معك حق، بغيّرها." and by[k]["scene"] == 8
    sc2 = P.parse(SAMPLE2)
    c = CMP.compare(sc, n, sc2, {i["id"]: i for i in items})
    assert next(x for x in c["previous_notes"] if x["id"] == k)["writer"]["replies"] == ["معك حق، بغيّرها."]


def test_cards_board_table_and_reorder(tmp_path, sc):
    stats = D.digest(sc, "short-one-moment")
    page = CARDS.board(sc, stats, full_notes(), "ar")
    assert page.count('class="c"') == 13 and "الضغط" in page
    path = tmp_path / "cards.docx"
    CARDS.table(sc, full_notes(), "ar").save(path)
    from docx import Document
    d = Document(path)
    for row in d.tables[0].rows[1:]:
        n = int(row.cells[0].text)
        row.cells[1].paragraphs[0].add_run({6: "7", 7: "6", 9: "x"}.get(n, str(n)))
    d.save(path)
    order, dropped = CARDS.read_order(path, 13)
    assert order[5:7] == [7, 6] and dropped == [9] and len(order) == 12
    new = CARDS.reorder(sc, order)
    assert len(new["scenes"]) == 12 and new["scenes"][5]["was"] == 7 and new["scenes"][5]["place"] == "البحر المفتوح"


def test_arc_chart(sc):
    n = full_notes()
    n["arcs"] = {"مريم": {"start": "تنتظر", "end": "تبحر", "points": {"1": {"value": -1}, "13": {"value": 2, "want": "أن تبحر"}}}}
    assert A.check(sc, n)[0] == []
    html = C.arc_chart(D.digest(sc), n, "ar")
    assert "أقواس الشخصيات" in html and "تنتظر" in html and "#2a78d6" in html
    n["arcs"]["مريم"]["points"]["99"] = {"value": 0}
    assert any("unknown scenes" in e for e in A.check(sc, n)[0])


# ---- v0.4: yardstick, quality checks, PDF/FDX, focused passes, tension overlay ----
def test_every_lens_has_anchors():
    for L in A.LENS.values():
        assert set(L["anchors"]) == {"green", "amber", "red"}
        assert all(L["anchors"][v]["ar"] and L["anchors"][v]["en"] for v in L["anchors"])


def test_quality_warnings():
    n = notes()
    n["lenses"][0]["summary"] = "جملة. " * 5
    n["notes"] += [{"scene": 2, "lens": "dialogue", "kind": "issue", "text": "الحوار يحتاج إلى عمق."},
                   {"scene": 2, "lens": "dialogue", "kind": "issue", "text": "الحوار يحتاج إلى عمق!"}]
    w = A.quality(n)
    assert any("summary" in x for x in w) and any("end with a question" in x for x in w)
    assert any("generic phrase" in x for x in w) and any("nearly the same" in x for x in w)


def test_scope_first_pages_and_lenses(sc):
    n = notes(lenses=[{"id": i, "verdict": "green", "summary": "x"} for i in ("premise", "character", "dialogue", "visual", "pacing")])
    n["scope"] = {"kind": "first-pages", "pages": 3}
    assert A.check(sc, n)[0] == [] and "أول ٣ صفحات" in A.scope_label(n, "ar")
    n2 = notes(lenses=[{"id": "dialogue", "verdict": "green", "summary": "x"}], notes=[{"scene": 2, "lens": "dialogue", "kind": "strength", "text": "x"}])
    n2["scope"] = {"kind": "lenses", "lenses": ["dialogue"]}
    assert A.check(sc, n2)[0] == []
    n2["scope"] = {"kind": "full"}
    assert any("has no verdict" in e for e in A.check(sc, n2)[0])


@pytest.mark.skipif(not SOFFICE or not shutil.which("pdftotext"), reason="needs LibreOffice and pdftotext")
def test_arabic_pdf_round_trip(sc, tmp_path):
    import re as _re
    out = tmp_path / "f.docx"
    F.build_docx(sc).save(out)
    pdf, err = S.to_pdf(out, tmp_path)
    assert pdf, err
    back = P.parse(pdf)
    assert back["title_page"]["title"] == "المرساة" and back["title_page"]["draft_date"] == "٢٠٢٦/١٠/٢٥"
    assert [e["type"] for e in back["elements"]] == [e["type"] for e in sc["elements"]]
    norm = lambda t: _re.sub(r"\W", "", _re.sub(r"^\s*(?:مشهد\s*)?[\d٠-٩]+\s*[.\-–]?\s*", "", t))  # noqa: E731
    assert all(norm(a["text"]) == norm(b["text"]) for a, b in zip(back["elements"], sc["elements"]))


def test_fdx_round_trip(tmp_path):
    f = tmp_path / "x.fountain"
    f.write_text("INT. KITCHEN - NIGHT\n\nA kettle.\n\nSARA\n(quietly)\nNot tonight.\n\nCUT TO:\n\nEXT. STREET - DAY\n\nRain.\n", encoding="utf-8")
    s = P.parse(f)
    fdx = tmp_path / "x.fdx"
    fdx.write_text(F.to_fdx(s), encoding="utf-8")
    back = P.parse(fdx)
    assert [e["type"] for e in back["elements"]] == [e["type"] for e in s["elements"]]
    assert back["scenes"][1]["int_ext"].startswith("EXT")


def test_show_first_pages(sc, tmp_path):
    d = D.digest(sc)
    d["pages"] = 6
    st = tmp_path / "st.json"
    st.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
    import show as SH
    keep = SH.wanted(["--pages", "2", "--stats", str(st)], 13)
    assert 1 in keep and 13 not in keep


def test_tension_overlay(sc):
    d = D.digest(sc)
    n_old = full_notes()
    n_new = full_notes()
    n_new["tension"] = {str(i): 3 for i in range(1, 14)}
    html = C.tension_curve(d, n_new, None, "ar", (d, n_old))
    assert "المسودة السابقة" in html and "stroke-dasharray=\"6 5\"" in html


def test_report_tables_do_not_clash_with_scene_strip(sc):
    import report as R
    page = R.build(sc, full_notes(), D.digest(sc, "short-one-moment"), "ar")
    assert '<table class="strip"' not in page and page.count('class="vtable"') == 2  # voices + stripboard


def test_first_pages_charts_only_show_read_scenes(sc):
    d = D.digest(sc)
    d["pages"] = 6
    n = full_notes()
    n["scope"] = {"kind": "first-pages", "pages": 3}
    st = C.in_scope(d, n)
    assert 0 < len(st["scene_rows"]) < 13 and all(int(k) <= max(r["n"] for r in st["scene_rows"]) for row in st["presence_grid"].values() for k in row)


# ---- story idea → template ----
import story_template as ST  # noqa: E402

MAP = json.loads((ROOT / "tests" / "fixtures" / "story_map_al_mirsah.json").read_text(encoding="utf-8"))


def test_story_map_checks_quotes_are_the_writers_words():
    assert ST.check(MAP)[0] == []
    bad = json.loads(json.dumps(MAP))
    bad["beats"]["turn"]["fragments"] = ["العم يُصاب في عاصفة"]  # an invented event
    bad["heard"][0]["open_question"] = "اكتب مشهداً عن البحر."  # not a question
    errs, warns = ST.check(bad)
    assert any("not the writer's words" in e for e in errs) and any("not a question" in w for w in warns)
    full = json.loads(json.dumps(MAP))
    full["beats"]["turn"]["fragments"] = ["عمها يرفض"]
    assert any("no gaps" in w for w in ST.check(full)[1])


def test_scene_counts_follow_running_time():
    st = next(s for s in S.STRUCTURES["structures"] if s["id"] == "short-one-moment")
    c = ST.scene_counts(st, 12, "short")
    assert c["pressure"] > c["hook"] and 8 <= sum(c.values()) <= 12


def test_story_template_docx_and_map(tmp_path):
    doc = ST.build_docx(MAP)
    out = tmp_path / "t.docx"
    doc.save(out)
    from docx import Document
    texts = [p.text for p in Document(out).paragraphs]
    assert MAP["story"] in texts  # the story, untouched
    assert "قصتك لا تقول هذا بعد." in texts  # the turn is a gap
    assert "«عمها يرفض»" in texts
    page = ST.build_map(MAP)
    assert 'class="gap"' in page and page.count('class="frag"') == 3 and 'dir="rtl"' in page


def test_story_working_title_and_duplicate_fragments():
    m = json.loads(json.dumps(MAP))
    m.pop("title")
    assert ST.working_title(m) == "بنت عمرها ١٧"
    m["story"] = "اممم فكرة فيلم يعني... رجل كبير في السن"
    assert ST.working_title(m) == "رجل كبير في"
    d = json.loads(json.dumps(MAP))
    d["beats"]["turn"]["fragments"] = ["عمها يرفض"]
    assert any("two beats" in w for w in ST.check(d)[1])


# --- documentary: structures, proposal template, story map ---
import _doc as DOC  # noqa: E402

DOC_MAP = json.loads((ROOT / "tests" / "fixtures" / "story_map_sur_doc.json").read_text(encoding="utf-8"))


def test_documentary_structures_and_proposal():
    docs = [s for s in S.STRUCTURES["structures"] if s["fits"] == ["documentary"]]
    assert len(docs) == 7 and all(s["id"].startswith("doc-") and DOC.is_doc(s["id"]) for s in docs)
    assert not DOC.is_doc("three-act")
    ids = [x["id"] for x in DOC.PROPOSAL]
    assert {"access", "approach", "archive", "why_now"} <= set(ids) and "consent" not in ids
    assert all(x["q"] and x["q_ar"] and x["name_ar"] for x in DOC.PROPOSAL)


def test_documentary_blank_template(tmp_path):
    from docx import Document
    out = tmp_path / "d.docx"
    B.build("doc-character", 25, "ar").save(out)
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    assert "الوصول" in texts and "الأساسات" not in texts  # proposal sections instead of fiction foundations
    assert any("الدقائق" in c.text for t in d.tables for c in t.rows[0].cells)  # timed in minutes
    av = d.tables[-1]
    assert [c.text for c in av.rows[0].cells] == ["#", "الصورة", "الصوت"] and len(av.rows) == 9
    assert av._tbl.tblPr.find(S.qn("w:bidiVisual")) is not None
    assert not any(p.style.name.startswith("Scene") for p in d.paragraphs)  # no screenplay samples


def test_documentary_story_map(tmp_path):
    assert ST.check(DOC_MAP)[0] == []
    bad = json.loads(json.dumps(DOC_MAP))
    bad["sections"]["archive"]["fragments"] = ["أرشيف التلفزيون العماني"]  # not in the idea
    bad["sections"]["consent"] = {}
    errs = ST.check(bad)[0]
    assert any("not the writer's words" in e for e in errs) and any("not a proposal section" in e for e in errs)
    wrong = json.loads(json.dumps(DOC_MAP))
    wrong["structure"] = "three-act"
    assert any("does not match" in e for e in ST.check(wrong)[0])
    noacc = json.loads(json.dumps(DOC_MAP))
    noacc["heard"] = [h for h in noacc["heard"] if h["element"] != "access"]
    assert any("access" in w for w in ST.check(noacc)[1])
    assert ST.working_title(DOC_MAP) == "آخر صانعي السفن"
    out = tmp_path / "m.docx"
    ST.build_docx(DOC_MAP).save(out)
    from docx import Document
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    assert "«صور قديمة لجد سالم من أرشيف عائلي»" in texts and "الأشخاص الذين ذكرتهم" in texts
    assert [c.text for c in d.tables[-1].rows[0].cells] == ["#", "الصورة", "الصوت"]
    page = ST.build_map(DOC_MAP)
    assert "المقترح" in page and "لم تذكرها قصتك بعد" in page and "دقيقة" in page


# --- documentary: transcripts → bites → paper edit ---
import paper_edit as PE  # noqa: E402
import transcripts as TR  # noqa: E402

DOCS = ROOT / "samples" / "script-desk" / "doc"


def _bites():
    return TR.ingest([DOCS / "salem-interview.srt", DOCS / "workshop-day.txt"], interviewer="المخرج")


def test_transcripts_srt_and_txt():
    data = _bites()
    by = {b["id"]: b for b in data["bites"]}
    assert [f["letter"] for f in data["files"]] == ["A", "B"]
    assert by["A002"]["text"].endswith("أبغى أروح البحر مع العيال.")  # two cues of one sentence joined
    assert by["A002"]["tc_in"] == "00:00:09" and by["A002"]["tc_out"] == "00:00:17"
    assert by["A005"]["direction"] == ["يؤشر على راسه"] and "(" not in by["A005"]["text"]  # direction kept apart
    assert by["A001"].get("interviewer") and not by["A003"].get("interviewer")
    assert by["B004"]["speaker"] == "رحيم" and by["B004"]["in"] == 20
    assert by["B009"]["text"] == "" and by["B009"]["direction"]  # a silence is a moment too
    assert by["B008"]["out"] == 72  # out point = next line's in


def test_transcripts_other_formats(tmp_path):
    vtt = tmp_path / "a.vtt"
    vtt.write_text("WEBVTT\n\n00:01.000 --> 00:03.500\n<v Salem>Speaker 1: The wood tells you\n\n00:03.600 --> 00:05.000\nwhere it wants to bend.\n", encoding="utf-8")
    txt = tmp_path / "b.txt"
    txt.write_text("Speaker 1  00:00:05\nWe start at dawn.\nBefore the heat.\n\nSpeaker 2 (00:00:12): And you?\n", encoding="utf-8")
    from docx import Document
    d = Document()
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([("00:00:01", "سالم", "بسم الله"), ("٠٠:٠٠:٠٤", "سالم", "نبدأ."), ("00:00:09", "رحيم", "زين")]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    dx = tmp_path / "c.docx"
    d.save(dx)
    data = TR.ingest([vtt, txt, dx], names={"Speaker 1": "Salem"})
    b = data["bites"]
    assert b[0]["speaker"] == "Salem" and b[0]["text"] == "The wood tells you where it wants to bend."
    txtb = [x for x in b if x["file"] == "b.txt"]
    assert txtb[0]["text"] == "We start at dawn. Before the heat." and txtb[0]["speaker"] == "Salem" and txtb[1]["speaker"] == "Speaker 2"
    dxb = [x for x in b if x["file"] == "c.docx"]
    assert [x["speaker"] for x in dxb] == ["سالم", "سالم", "رحيم"] and dxb[1]["in"] == 4
    assert TR.expand(["A001-A003", "B002"], _bites()["bites"]) == ["A001", "A002", "A003", "B002"]


def test_paper_edit(tmp_path):
    bites = _bites()["bites"]
    sel = json.loads((ROOT / "tests" / "fixtures" / "selects_sur_doc.json").read_text(encoding="utf-8"))
    errs, warns = PE.check(sel, bites)
    assert errs == [] and any("A012→A016" in w for w in warns) and not any("A006→A009" in w for w in warns)
    bad = json.loads(json.dumps(sel))
    bad["sections"][0]["items"] += [{"bite": "A099"}, {"bite": "A001"}, {"bite": "A005", "sound": "typed words"}, {"bite": "A006", "text": "x"}]
    errs, warns = PE.check(bad, bites)
    assert any("no bite A099" in e for e in errs) and any("remove 'sound'" in e for e in errs) and any("unknown fields" in e for e in errs)
    assert any("interviewer" in w for w in warns) and any("used twice" in w for w in warns)
    out = tmp_path / "p.docx"
    PE.build_docx(sel, bites).save(out)
    from docx import Document
    d = Document(out)
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    by = {b["id"]: b for b in bites}
    assert any(f"«{by['A016']['text']}»" in c for c in cells)  # verbatim from the transcript
    assert any("ينقصنا" in c for c in cells) and any("✂" in c for c in cells)
    assert any(c == "0:54" for c in cells)  # total running time
    PE.write_csv(sel, bites, tmp_path / "p.csv")
    rows = (tmp_path / "p.csv").read_text(encoding="utf-8-sig").splitlines()
    assert rows[0].startswith("order,section,bite") and len(rows) == 12 and "GAP:" in rows[8]


# --- documentary: two-column A/V script ---
import av_script as AV  # noqa: E402


def test_av_script_text_docx_and_reader(tmp_path):
    av = AV.read(DOCS / "sur-av-draft.txt", "آخر صانعي السفن")
    assert av["lang"] == "ar" and len(av["rows"]) == 16 and len(AV.sections(av)) == 5
    r = av["rows"][3]
    assert r["section"] == "العالم: سفينة في الرأس" and r["sound"].startswith("سالم: الرسم")
    words, sec = AV.spoken(av)
    assert 80 < words < 110 and sec > 30  # ambience («موج خفيف») is not counted as speech
    out = tmp_path / "av.docx"
    AV.build_docx(av).save(out)
    back = AV.read(out)  # its own Word file reads back the same
    assert [(x["section"], x["picture"], x["sound"]) for x in back["rows"]] == [(x["section"], x["picture"], x["sound"]) for x in av["rows"]]
    from docx import Document
    nums = [t.rows[-1].cells[0].text for t in Document(out).tables]
    assert nums[-1] == "١٦"  # numbered through the whole script
    sc = AV.to_script(av)
    assert sc["schema"] == "zamakan-script/1" and len(sc["scenes"]) == 5
    names = {e["name"] for e in sc["elements"] if e["type"] == "character"}
    assert {"سالم", "تعليق صوتي", "الابن", "الصوت"} <= names
    txt = tmp_path / "en.txt"
    txt.write_text("SEQUENCE 1: Dawn\nVIDEO: Harbour at night.\nAUDIO: Waves.\nVIDEO: Salem's hands.\nAUDIO: SALEM: The wood tells you.\n", encoding="utf-8")
    en = AV.read(txt)
    assert en["lang"] == "en" and en["rows"][0]["section"] == "1: Dawn" and en["rows"][1]["sound"] == "SALEM: The wood tells you."


# --- documentary reader: doc viewpoints on an A/V script, a paper edit and a treatment ---
import doc_to_script as D2S  # noqa: E402

DOC_NOTES = json.loads((ROOT / "tests" / "fixtures" / "notes_sur_av.json").read_text(encoding="utf-8"))


@pytest.fixture
def film_lenses_after():
    yield
    A.use_lenses("film")  # other tests expect the film viewpoints


def test_doc_lenses_file():
    doc = json.loads((S.SKILL / "references" / "lenses_doc.json").read_text(encoding="utf-8"))
    film = json.loads((S.SKILL / "references" / "lenses.json").read_text(encoding="utf-8"))
    assert len(doc["lenses"]) == 10 and doc["verdicts"] == film["verdicts"] and doc["families"] == film["families"]
    for L in doc["lenses"]:
        assert L["hex"] == doc["families"][L["family"]]["hex"] and set(L["anchors"]) == {"green", "amber", "red"} and L["name_ar"]
    assert {"doc_ethics", "doc_access", "doc_archive", "doc_evidence"} <= {L["id"] for L in doc["lenses"]}
    assert set(doc["first_pages"]) <= {L["id"] for L in doc["lenses"]}


def test_doc_reader_on_av_script(tmp_path, film_lenses_after):
    sc = AV.to_script(AV.read(DOCS / "sur-av-draft.txt", "آخر صانعي السفن"))
    errs, warns = A.check(sc, DOC_NOTES)
    assert errs == [] and "doc_access" in A.LENS and "premise" not in A.LENS
    missing = json.loads(json.dumps(DOC_NOTES))
    missing["lenses"] = [x for x in missing["lenses"] if x["id"] != "doc_archive"]
    assert any("doc_archive" in e for e in A.check(sc, missing)[0])
    film = json.loads(json.dumps(DOC_NOTES))
    film.pop("lens_set")  # film viewpoints: doc ids are unknown there
    assert any("unknown lens" in e for e in A.check(sc, film)[0])
    d = D.digest(sc, "doc-character")
    dd = d["doc"]
    assert any("آخر صانع" in x for x in dd["superlative"]) and any("إعادة تمثيل" in x for x in dd["reconstruction"])
    assert len(dd["narration"]) == 2 and len(dd["archive"]) == 2 and any("عامل" in x for x in dd["unnamed_people"])
    assert not any("مثل أول" in x for x in dd["superlative"])  # colloquial «أول» is not a claim
    shown = D.show(d)
    assert "documentary (av)" in shown and "pages" not in shown.split("\n")[0] and "protagonist" not in shown
    A.check(sc, DOC_NOTES)
    doc = A.annotate(sc, DOC_NOTES, "ar")
    out = tmp_path / "n.docx"
    doc.save(out)
    import zipfile
    z = zipfile.ZipFile(out)
    comments = z.read("word/comments.xml").decode("utf-8")
    assert comments.count("<w:comment ") == len(DOC_NOTES["notes"]) and "الدقة والأدلة" in comments
    body = z.read("word/document.xml").decode("utf-8")
    first_tbl = body.index("<w:tbl>", body.index("الافتتاح"))
    assert body.index("commentRangeStart", first_tbl) < body.index("</w:tbl>", first_tbl)  # comments sit inside the A/V table
    import report as R
    page = R.build(sc, DOC_NOTES, d, "ar")
    assert "زمن الكلام" in page and "لوحة التصوير" not in page


def test_doc_to_script_paper_edit_and_treatment(tmp_path):
    bites = _bites()["bites"]
    sel = json.loads((ROOT / "tests" / "fixtures" / "selects_sur_doc.json").read_text(encoding="utf-8"))
    sc = PE.to_script(sel, bites)
    assert sc["kind"] == "documentary" and len(sc["scenes"]) == 3
    lines = [e for e in sc["elements"] if e["type"] == "dialogue" and e.get("bite") == "A016"]
    assert lines and lines[0]["text"] == {b["id"]: b for b in bites}["A016"]["text"]
    assert any(e["text"].startswith("[ينقصنا") for e in sc["elements"])
    tr = tmp_path / "t.md"
    tr.write_text("# صانعو السفن\nمقدمة قصيرة عن الفيلم.\n## الوصول\nسالم وافق.\nالابن لم يرد.\n## الأرشيف\nصور العائلة.\n", encoding="utf-8")
    t = D2S.treatment(tr)
    assert t["title_page"]["title"] == "صانعو السفن" and [s["heading"] for s in t["scenes"]] == ["المقدمة", "الوصول", "الأرشيف"]
    assert [e["text"] for e in t["elements"] if e.get("scene") == 2 and e["type"] == "action"] == ["سالم وافق.", "الابن لم يرد."]
    from av_script import script_docx
    pm = {}
    script_docx(t, pm)
    assert all(i in pm for i, e in enumerate(t["elements"]))  # every line can carry a comment


# --- documentary helpers: archive log, shoot list ---
import doc_helpers as DH  # noqa: E402


def test_archive_log_and_merge(tmp_path):
    sc = AV.to_script(AV.read(DOCS / "sur-av-draft.txt", "آخر صانعي السفن"))
    items = DH.archive(sc)
    assert [i["id"] for i in items] == ["AR01", "AR02"] and "صف 9" in items[0]["used_in"] and items[1]["rights"] == "للسؤال"
    p = tmp_path / "log.csv"
    items[0].update(owner="عائلة سالم", rights="مُصرّح بها")
    items.append({"id": "AR03", "item": "أغنية قديمة", "used_in": "x", "rights": "طُلبت"})
    DH.write_csv(items, DH.ARCHIVE_COLS, p)
    again = DH.archive(sc, p)
    assert again[0]["owner"] == "عائلة سالم" and again[0]["rights"] == "مُصرّح بها"  # filled-in status survives
    assert again[2]["item"] == "أغنية قديمة" and "no longer in this draft" in again[2]["notes"]  # nothing silently dropped
    DH.archive_docx(again, "ar", "x").save(tmp_path / "a.docx")


def test_shoot_list(tmp_path):
    av = AV.read(DOCS / "sur-av-draft.txt", "آخر صانعي السفن")
    plan = DH.skeleton(av)
    assert len(plan["items"]) == 16 and all(i["access"] == "pending" for i in plan["items"])
    plan["items"] = plan["items"][:3]
    plan["items"][0]["access"] = "confirmed"
    plan["items"][2]["access"] = "no"
    assert DH.check_plan(plan) == []
    doc, items = DH.shoots_docx(plan)
    assert [i["access"] for i in items] == ["pending", "no", "confirmed"]  # waiting on a yes comes first
    out = tmp_path / "s.docx"
    doc.save(out)
    from docx import Document
    assert "! بانتظار الموافقة" in [c.text for c in Document(out).tables[0].rows[1].cells]
    assert DH.check_plan({"items": [{"what": "", "access": "maybe"}]})
