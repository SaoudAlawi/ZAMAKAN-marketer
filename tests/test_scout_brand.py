"""Scout outputs dressed in the client's own brand (feature-1 design system) + the beauty/brand check."""
import json
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / ".claude/skills/zamakan-reference-scout/scripts"
FIX = REPO / "evals/fixtures/scout"
sys.path.insert(0, str(SCRIPTS))

import brandkit as K  # noqa: E402
import build_outputs as B  # noqa: E402
import check_outputs as C  # noqa: E402


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    """clients/rawaq/design-system built from the gold spec, and a scout job folder under it."""
    root = tmp_path_factory.mktemp("agency")
    ds = root / "clients/rawaq/design-system"
    subprocess.run([sys.executable, REPO / ".claude/skills/zamakan-design-system/scripts/build_design_system.py",
                    REPO / "evals/fixtures/rawaq/expected-spec.json", "--input", REPO / "samples/rawaq/input",
                    "--out", ds, "--no-preview"], check=True, capture_output=True)
    run = json.loads((FIX / "rawaq-ramadan-sample-run.json").read_text(encoding="utf-8"))
    for r in run["references"]:
        r.pop("thumbnail", None)  # offline: no downloads
    job = root / "clients/rawaq/jobs/2026-10-13-scout-ramadan"
    job.mkdir(parents=True)
    (job / "references.json").write_text(json.dumps(run, ensure_ascii=False), encoding="utf-8")
    sys.argv = ["build_outputs.py", str(job / "references.json"), "--out", str(job), "--formats", "board,deck,report",
                "--brief", str(FIX / "rawaq-ramadan-brief.json"), "--no-pdf"]
    B.main()
    return ds, job


def test_kit_reads_the_whole_design_system(client):
    ds, job = client
    kit = K.load_kit(None, job)
    th = kit["theme"]
    assert kit["source"] == "client" and kit["client_name"] == "Rawaq Coffee Roasters"
    assert th["bg"] == "#F3E7D3" and th["dark"] == "#4A2C21" and th["on_dark"] == "#F3E7D3"  # sand / date brown, by role
    assert th["display_font"] == "Playfair Display" and th["ar_display_font"] == "Reem Kufi"
    # logos: the variant approved for the background it sits on
    assert kit["logos"]["on_dark"]["id"] == "reversed" and kit["logos"]["on_light"]["id"] == "horizontal"
    assert kit["logos"]["mark_light"]["id"] == "symbol"
    assert {f["family"] for f in kit["font_files"]} == {"Playfair Display", "IBM Plex Sans", "Reem Kufi", "IBM Plex Sans Arabic"}
    assert kit["radius"] == "0px" and kit["arch"] and kit["divider"] == "1px solid #E3A33B" and kit["arabic_leads"]
    assert "Mashrabiya pattern tile (SVG)" in kit["missing_assets"]
    # saffron is never text on sand (brand rule), the readable accent is used instead
    assert th["accent_text"] != "#E3A33B"


def test_board_carries_brand_assets(client):
    _, job = client
    h = (job / "board.html").read_text(encoding="utf-8")
    for fam in ("Playfair Display", "IBM Plex Sans", "Reem Kufi", "IBM Plex Sans Arabic"):
        assert f"@font-face{{font-family:'{fam}'" in h  # the client's own files, embedded
    assert "fonts.googleapis.com" not in h  # works offline
    assert 'class="topbar"><img src="data:image/svg+xml;base64,' in h  # reversed logo on the cover
    assert "--r:0px" in h and "--dark:#4A2C21" in h and 'class="divider"' in h
    assert "Styled with the Rawaq Coffee Roasters design system" in h


def test_deck_carries_brand_assets(client):
    from pptx import Presentation
    _, job = client
    prs = Presentation(job / "deck.pptx")
    pics = lambda s: [sh for sh in s.shapes if sh.shape_type == 13]  # noqa: E731
    assert len(pics(prs.slides[0])) >= 1  # reversed logo (no thumbnails offline)
    assert all(len(pics(s)) >= 1 for s in list(prs.slides)[1:-1])  # brand mark on every content slide
    assert str(prs.slides[0].background.fill.fore_color.rgb) == "4A2C21"
    fonts = {r.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert {"Playfair Display", "IBM Plex Sans", "Reem Kufi"} <= fonts  # Arabic territory names are headings: display face
    theme_xml = next(r.target_part.blob for r in prs.slide_masters[0].part.rels.values() if r.reltype.endswith("/theme")).decode()
    assert '<a:hlink><a:srgbClr val="4A2C21"/>' in theme_xml  # links in the brand colour, not default blue


def test_beauty_check_passes_on_brand_build(client):
    _, job = client
    out = C.run_checks(job)
    assert out["ok"], out["errors"]
    assert set(out["info"]["board"]["desktop"]["fonts_loaded"]) >= {"Playfair Display", "IBM Plex Sans", "Reem Kufi"}
    assert (job / "check/board-desktop.png").exists() and (job / "check/board-phone.png").exists()


def test_beauty_check_catches_ugly(client, tmp_path):
    _, job = client
    bad = tmp_path / "bad"
    bad.mkdir()
    h = (job / "board.html").read_text(encoding="utf-8")
    h = h.replace("</style>", ".ins p{color:#E3A33B!important} .ar-d{letter-spacing:.2em}</style>", 1)  # saffron on sand; tracked Arabic
    h = h.replace("<main>", '<main><div style="width:1900px;height:4px"></div>', 1)  # sideways scroll
    (bad / "board.html").write_text(h, encoding="utf-8")
    out = C.run_checks(bad)
    errs = " ".join(out["errors"])
    assert not out["ok"] and "low contrast" in errs and "scrolls sideways" in errs and "letter-spaced" in errs


def test_house_style_with_client_logo(tmp_path):
    """A client with no design system yet: house look, but the client's own logo (on a white panel on the dark cover)."""
    logo = next((REPO / "samples/rawaq/input").rglob("*.png"))
    kit = K.with_client_logo(K.load_kit("zamakan"), logo, tmp_path, "Rawaq")
    assert kit["source"] == "house" and kit["client_name"] == "Rawaq"
    assert kit["logos"]["on_light"]["png"] == str(logo)
    panel = Path(kit["logos"]["on_dark"]["png"])
    from PIL import Image
    im = Image.open(panel)
    assert panel.exists() and im.getpixel((im.width // 2, 2))[:3] == (255, 255, 255)


def test_dark_cover_gets_logo_on_white_panel(tmp_path):
    """No reversed logo approved on the dark cover: the approved light logo goes on a white panel (never redrawn)."""
    from PIL import Image
    logo = next((REPO / "samples/rawaq/input").rglob("*.png"))
    panel = K.logo_panel(logo, tmp_path / "p.png")
    im = Image.open(panel)
    assert im.getpixel((im.width // 2, 1))[:3] == (255, 255, 255) and im.width > Image.open(logo).width


def test_labels_on_brand_fills_stay_readable():
    """A mid-blue brand colour with white labels (4.3:1) is nudged just enough to pass 4.5:1; rings too."""
    import theme as T
    th = T.resolve({"bg": "#FFFFFF", "ink": "#0A3550", "accent": "#0D80BF"}, ["#0D80BF", "#99BF3C", "#DEDEDB"])
    assert T.contrast(th["on_accent"], th["accent"]) >= 4.5
    for k, c in th["ring"].items():
        assert T.contrast(th["ring_on"][k], c) >= 4.5, k
    assert min(T.contrast(th["muted"], c) for c in (th["bg"], th["panel"], th["chip"])) >= 4.5


def test_deck_text_fitting_measures_wide_fonts():
    import render_deck as RD
    house = K.load_kit("zamakan")
    assert RD.avg_char_width(house) == 0.52
    wide = dict(house, font_files=[{"path": REPO / ".claude/skills/zamakan-reference-scout/assets/fonts/IBMPlexSans[wdth,wght].ttf", "script": "latin"}])
    assert abs(RD.avg_char_width(wide) - 0.52) < 0.01  # the font the default was tuned on
