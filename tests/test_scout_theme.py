"""DESIGN.md theming of scout boards and decks: house style, client design system, awesome-design-md brands."""
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / ".claude/skills/zamakan-reference-scout/scripts"
FIX = REPO / "evals/fixtures/scout"
sys.path.insert(0, str(SCRIPTS))

import build_outputs as B  # noqa: E402
import theme as T  # noqa: E402
from test_scout_outputs import run_path  # noqa: E402,F401  (fixture)

PROSE_ONLY = """# Acme Coffee

## Colors
- Canvas white `#FFFFFF`, ink `#111111`, brand green `#00704A`, warm cream `#F2F0EB`.

## Typography
### Font Family
- **Display**: `Lander Tall`, Georgia, serif
- **Primary:** `SoDoSans, "Helvetica Neue", Arial, sans-serif`
- **Loading Fallback:** `Arial`
"""

BROKEN_YAML = """---
version: alpha
name: Broken
description: A system: unquoted colon breaks YAML here: badly
colors:
  primary: "#E60023"
  canvas: "#FFFFFF"
  ink: "#111111"
---
# Broken
"""


def readable(th):
    c = T.contrast
    return all([c(th["ink"], th["bg"]) >= 4.5, c(th["ink"], th["panel"]) >= 4.5, c(th["muted"], th["panel"]) >= 4.5,
                c(th["accent_text"], th["bg"]) >= 4.5, c(th["on_dark"], th["dark"]) >= 4.5])


def test_house_theme():
    th = T.load("zamakan")
    assert th["name"] == "ZAMAKAN" and th["bg"] == "#F6F6F2" and th["ink"] == "#14213D"
    assert th["accent"] == "#0B7C7C"  # near-black primary (navy) is ink, the teal tertiary is the accent
    assert th["display_font"] == "Instrument Serif" and th["ar_display_font"] == "Reem Kufi" and th["display_weight"] == 400
    assert readable(th)


def test_prose_only_design_md(tmp_path):
    p = tmp_path / "DESIGN.md"
    p.write_text(PROSE_ONLY, encoding="utf-8")
    th = T.load(str(p))
    assert th["display_font"] == "Lander Tall" and th["body_font"] == "SoDoSans"
    assert th["accent"] == "#00704A" and th["bg"] == "#FFFFFF" and readable(th)


def test_malformed_front_matter_is_recovered(tmp_path):
    p = tmp_path / "DESIGN.md"
    p.write_text(BROKEN_YAML, encoding="utf-8")
    th = T.load(str(p))
    assert th["name"] == "Broken" and th["accent"] == "#E60023" and readable(th)


@pytest.mark.parametrize("colors", [
    {"canvas": "#000000", "ink": "#FFFFFF", "primary": "#0066B1"},          # dark canvas, dim accent
    {"canvas": "#7A8ABA", "ink": "#21242E", "primary": "#E60012"},          # mid-tone canvas
    {"background": "#FFFFFF", "on-surface": "#222222", "primary": "#FFD400"},  # yellow accent on white
])
def test_contrast_is_enforced(colors):
    th = T.resolve(colors, list(colors.values()))
    assert readable(th)
    assert th["accent"] == colors["primary"]  # fills keep the brand colour; only accent_text is nudged


def test_client_design_system_is_auto_detected(run_path, tmp_path):
    ds = tmp_path / "clients/rawaq/design-system"
    subprocess.run([sys.executable, REPO / ".claude/skills/zamakan-design-system/scripts/build_design_system.py",
                    REPO / "evals/fixtures/rawaq/expected-spec.json", "--input", REPO / "samples/rawaq/input",
                    "--out", ds, "--no-preview"], check=True, capture_output=True)
    out = tmp_path / "clients/rawaq/jobs/scout"
    th = T.choose(None, out)
    assert th["origin"].endswith("clients/rawaq/design-system/DESIGN.md") and th["name"] == "Rawaq Coffee Roasters"
    assert th["bg"] == "#F3E7D3" and readable(th)
    # an explicit --design wins over the client's own
    assert T.choose("zamakan", out)["name"] == "ZAMAKAN"


def test_deck_uses_theme(run_path, tmp_path):
    from pptx import Presentation
    th = T.load("zamakan")
    out = tmp_path / "job"
    sys.argv = ["build_outputs.py", str(run_path), "--out", str(out), "--formats", "board,deck", "--no-pdf", "--design", "zamakan"]
    B.main()
    prs = Presentation(out / "deck.pptx")
    cover, brief = prs.slides[0], prs.slides[1]
    assert str(cover.background.fill.fore_color.rgb) == th["dark"].lstrip("#")
    assert str(brief.background.fill.fore_color.rgb) == th["bg"].lstrip("#")
    fonts = {r.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs}
    assert {"Instrument Serif", "IBM Plex Sans", "Reem Kufi"} <= fonts
    board = (out / "board.html").read_text(encoding="utf-8")
    assert "--accent:#0B7C7C" in board and "@font-face{font-family:'Instrument Serif'" in board  # house fonts ship with the skill


def test_awesome_name_is_checked():
    with pytest.raises(ValueError):
        T.load("awesome:../../etc/passwd")


def _designmd_cli():
    for c in (os.environ.get("ZAMAKAN_DESIGNMD"), shutil.which("design.md")):
        if c and Path(c).exists():
            return c
    return None


@pytest.mark.skipif(not _designmd_cli(), reason="DESIGN.md CLI not installed (npm i -g @google/design.md)")
def test_house_design_md_passes_official_linter():
    r = subprocess.run([_designmd_cli(), "lint", str(T.HOUSE)], capture_output=True, text=True)
    report = json.loads(r.stdout)
    assert report["summary"]["errors"] == 0 and report["summary"]["warnings"] == 0, report


def _online():
    r = subprocess.run(["curl", "-sS", "-m", "10", "-o", "/dev/null", "-w", "%{http_code}", "https://www.example.org/"],
                       capture_output=True, text=True)
    return not os.environ.get("ZAMAKAN_OFFLINE") and r.stdout == "200"


@pytest.mark.skipif(not _online(), reason="no internet access")
def test_live_awesome_theme(tmp_path, monkeypatch):
    monkeypatch.setenv("ZAMAKAN_SCOUT_CACHE", str(tmp_path))
    th = T.load("awesome:nike")
    assert th["name"] == "Nike" and "awesome-design-md" in th["origin"] and readable(th)
    assert (tmp_path / "awesome-design-md/nike.md").exists()
